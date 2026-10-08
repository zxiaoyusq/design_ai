"""去重下载趋势 Excel 图片链接，并把同一文件路径回填到全部文章。"""

from __future__ import annotations

import argparse
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import UTC, datetime
import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile
from urllib.parse import urlsplit

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))
from prepare_trend_data import IMAGE_EXTENSIONS, PROJECT_ROOT, _image_metadata, download_image

DEFAULT_INPUT = PROJECT_ROOT / "ref/趋势信息0930.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data/trend_data/trend_info_0930"
DEFAULT_IMAGES = PROJECT_ROOT / "data/trend_data/images"
CHECKPOINT = "download-checkpoint.jsonl"


def digest(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def atomic_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, suffix=".tmp", delete=False) as stream:
            temporary = Path(stream.name)
            json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
            stream.write("\n")
        os.replace(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def read_articles(path: Path, sheet_name: str | None = None) -> tuple[list[dict], str]:
    """仅读文章身份、标题与 image_url；不读取已有图片解析字段。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            if len(workbook.sheetnames) != 1:
                raise ValueError("工作簿有多个工作表，请指定 --sheet")
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"找不到工作表：{sheet_name}")
        sheet = workbook[sheet_name]
        sheet.reset_dimensions()
        rows = iter(sheet.iter_rows(values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        if len([value for value in headers if value]) != len(set(filter(None, headers))):
            raise ValueError("表头包含重复字段")
        positions = {name: index for index, name in enumerate(headers) if name}
        missing = {"id", "image_url"} - set(positions)
        if missing:
            raise ValueError(f"缺少字段：{', '.join(sorted(missing))}")
        records, ids = [], set()
        for row_number, values in enumerate(rows, 2):
            if not any(value is not None and value != "" for value in values):
                continue
            raw_id = values[positions["id"]] if positions["id"] < len(values) else None
            if raw_id is None or not str(raw_id).strip():
                raise ValueError(f"第 {row_number} 行缺少 id")
            article_id = str(raw_id).strip()
            if article_id in ids:
                raise ValueError(f"第 {row_number} 行 id 重复：{article_id}")
            ids.add(article_id)
            def cell(field: str):
                index = positions.get(field)
                return values[index] if index is not None and index < len(values) else None
            raw_urls = cell("image_url")
            urls = [part.strip() for part in str(raw_urls or "").split("||") if part.strip()]
            records.append({
                "id": article_id,
                "source_row": row_number,
                "title": cell("title"),
                "title_zh": cell("title_zh"),
                "image_url": raw_urls,
                "images": [{"image_id": f"{article_id}_{i:03d}", "trend_id": article_id,
                            "index": i, "url": url, "local_path": None, "status": "pending"}
                           for i, url in enumerate(urls, 1)],
            })
        return records, sheet_name
    finally:
        workbook.close()


def existing_downloads(index_path: Path, images_root: Path, output_dir: Path) -> dict[str, dict]:
    """从旧索引按 URL 查找本地文件，只复用索引记录且 SHA-256 仍匹配的图片。"""
    if not index_path.is_file():
        return {}
    index = json.loads(index_path.read_text(encoding="utf-8"))
    image_root = images_root.resolve()
    result: dict[str, dict] = {}
    for article in index.get("trends", []):
        for image in article.get("images", []):
            url = image.get("url")
            relative = image.get("local_path")
            recorded_hash = image.get("sha256")
            if not isinstance(url, str) or not relative or image.get("status") != "downloaded":
                continue
            path = (index_path.parent / relative).resolve()
            try:
                path.relative_to(image_root)
                if not path.is_file() or not isinstance(recorded_hash, str) or digest(path) != recorded_hash:
                    continue
                metadata = _image_metadata(path)
            except (OSError, ValueError):
                continue
            result.setdefault(url, {**metadata, "path": path,
                                    "downloaded_at": image.get("downloaded_at"),
                                    "download_url": None, "resolved_url": None})
    return result


def groups(records: list[dict]) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for article in records:
        for image in article["images"]:
            grouped[image["url"]].append(image)
    return dict(grouped)


def cache_entry(url: str, refs: list[dict], cached: dict, output_dir: Path) -> dict:
    source = refs[0]
    path = cached["path"]
    try:
        relative = path.relative_to(output_dir.resolve())
    except ValueError:
        relative = Path(os.path.relpath(path, output_dir))
    return {**source, **{key: value for key, value in cached.items() if key != "path"},
            "url": url, "local_path": relative.as_posix(), "status": "downloaded",
            "cached": True, "reused_from_existing_index": True,
            "shared_reference_count": len(refs)}


def write_final(output: Path, records: list[dict], source: dict, unique_results: dict[str, dict]) -> dict:
    images = [image for article in records for image in article["images"]]
    successful = [image for image in images if image["status"] == "downloaded"]
    failures = [image for image in images if image["status"] == "failed"]
    counts = {
        "articles": len(records), "image_references": len(images),
        "unique_image_urls": len(unique_results),
        "downloaded_unique_urls": sum(item["status"] == "downloaded" for item in unique_results.values()),
        "failed_unique_urls": sum(item["status"] == "failed" for item in unique_results.values()),
        "references_with_local_files": len(successful), "references_unresolved": len(failures),
        "urls_reused_from_existing_index": sum(item.get("reused_from_existing_index", False) for item in unique_results.values()),
        "urls_reused_from_shared_store": sum(
            item.get("cached", False) and not item.get("reused_from_existing_index", False)
            for item in unique_results.values()),
    }
    payload = {"schema_version": "trend_image_inventory_v1", "generated_at": datetime.now(UTC).isoformat(),
               "source": source, "counts": counts, "trends": records}
    atomic_json(output / "trends.json", payload)
    temporary = output / "trends.jsonl.tmp"
    with temporary.open("w", encoding="utf-8") as stream:
        for article in records:
            stream.write(json.dumps(article, ensure_ascii=False, allow_nan=False) + "\n")
    os.replace(temporary, output / "trends.jsonl")
    atomic_json(output / "download_report.json", {
        **{key: value for key, value in payload.items() if key != "trends"},
        "unique_failures": [item for item in unique_results.values() if item["status"] == "failed"],
        "unresolved_references": failures,
    })
    return counts


def run(input_path: Path, output: Path, images_root: Path, *, sheet: str | None,
        workers: int, timeout: float, retries: int, dry_run: bool, resume: bool,
        retry_transient_failures: bool = False) -> dict:
    input_path, output, images_root = input_path.resolve(), output.resolve(), images_root.resolve()
    source_sha = digest(input_path)
    records, sheet_name = read_articles(input_path, sheet)
    by_url = groups(records)
    old_index = images_root.parent / "trends.json"
    cached = existing_downloads(old_index, images_root, output)
    already = set(by_url) & set(cached)
    plan = {"articles": len(records), "image_references": sum(map(len, by_url.values())),
            "unique_image_urls": len(by_url), "already_downloaded_urls": len(already),
            "new_unique_urls": len(by_url) - len(already), "sheet": sheet_name,
            "source_sha256": source_sha}
    print(json.dumps(plan, ensure_ascii=False), flush=True)
    if dry_run:
        return plan
    if workers < 1 or timeout <= 0 or retries < 0:
        raise ValueError("workers、timeout必须为正，retries不能小于0")
    if retry_transient_failures and not resume:
        raise ValueError("--retry-transient-failures 需要与 --resume 一起使用")

    checkpoint = output / CHECKPOINT
    manifest_path = output / "download_manifest.json"
    if resume:
        if not manifest_path.is_file() or not checkpoint.is_file():
            raise ValueError("找不到可续跑的下载清单/检查点")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("source_sha256") != source_sha or manifest.get("image_references") != plan["image_references"]:
            raise ValueError("源 Excel 或图片引用数量与中断运行不一致")
        unique_results = {}
        for line in checkpoint.read_text(encoding="utf-8").splitlines():
            if line.strip():
                entry = json.loads(line)
                unique_results[entry["url"]] = entry
        if not set(unique_results) <= set(by_url):
            raise ValueError("检查点含有不属于本次工作簿的 URL")
        if retry_transient_failures:
            # 仅重试可恢复的网络故障；404、策略拒绝及图片大小限制不重复请求。
            for url, result in list(unique_results.items()):
                error = str(result.get("error", ""))
                status = result.get("http_status")
                transient_http = status in (408, 429) or (isinstance(status, int) and status >= 500)
                transient_network = error.startswith(("TimeoutError:", "URLError:", "ConnectionError:"))
                if result.get("status") == "failed" and (transient_http or transient_network):
                    del unique_results[url]
    else:
        if output.exists():
            raise FileExistsError(f"输出目录已存在，若为中断任务请用 --resume：{output}")
        output.mkdir(parents=True)
        manifest = {"version": "1.0", "started_at": datetime.now(UTC).isoformat(),
                    "input": str(input_path), "source_sha256": source_sha, "sheet": sheet_name,
                    **{key: plan[key] for key in ("articles", "image_references", "unique_image_urls", "already_downloaded_urls", "new_unique_urls")},
                    "images_root": str(images_root), "workers": workers, "timeout_seconds": timeout, "retries": retries}
        atomic_json(manifest_path, manifest)
        checkpoint.touch()
        unique_results = {}

    for url in list(unique_results):
        saved = unique_results[url]
        if saved.get("status") == "downloaded":
            path = (output / saved["local_path"]).resolve()
            if not path.is_file() or digest(path) != saved.get("sha256"):
                del unique_results[url]

    pending = [url for url in by_url if url not in unique_results]
    for url in pending:
        if url in cached:
            unique_results[url] = cache_entry(url, by_url[url], cached[url], output)

    new_urls = [url for url in pending if url not in cached]
    done = 0
    with checkpoint.open("a", encoding="utf-8") as journal:
        with ThreadPoolExecutor(max_workers=workers) as pool:
            jobs = {pool.submit(download_image, {**by_url[url][0], "url": url}, output, timeout, retries,
                                images_dir=images_root): url for url in new_urls}
            for future in as_completed(jobs):
                url = jobs[future]
                result = future.result()
                result["shared_reference_count"] = len(by_url[url])
                unique_results[url] = result
                journal.write(json.dumps(result, ensure_ascii=False, allow_nan=False) + "\n")
                done += 1
                if done % 100 == 0 or done == len(new_urls):
                    journal.flush()
                    os.fsync(journal.fileno())
                    atomic_json(output / "progress.json", {**plan, "downloaded_unique_urls": sum(
                        value["status"] == "downloaded" for value in unique_results.values()),
                        "failed_unique_urls": sum(value["status"] == "failed" for value in unique_results.values()),
                        "new_urls_completed": done, "new_urls_total": len(new_urls),
                        "updated_at": datetime.now(UTC).isoformat()})
                    print(f"已处理新链接 {done}/{len(new_urls)}；索引复用 {len(already)} 个旧链接", flush=True)

    # 同 URL 在 Excel 内只下载一次，但保留每篇文章每个原始位置的对应关系。
    for url, refs in by_url.items():
        result = unique_results[url]
        for index, image in enumerate(refs):
            entry = {**result, "image_id": image["image_id"], "trend_id": image["trend_id"],
                     "index": image["index"], "url": url,
                     "shared_reference": index > 0 or url in cached,
                     "shared_reference_count": len(refs)}
            image.update(entry)
    counts = write_final(output, records, manifest, unique_results)
    atomic_json(output / "completion.json", {"completed_at": datetime.now(UTC).isoformat(), **counts})
    return counts


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--images-dir", type=Path, default=DEFAULT_IMAGES)
    parser.add_argument("--sheet")
    parser.add_argument("--workers", type=int, default=8)
    parser.add_argument("--timeout", type=float, default=30)
    parser.add_argument("--retries", type=int, default=2)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--resume", action="store_true", help="续跑本脚本保存了下载清单的中断任务")
    parser.add_argument("--retry-transient-failures", action="store_true",
                        help="续跑时只重试超时、连接故障及 408/429/5xx")
    args = parser.parse_args(argv)
    try:
        result = run(args.input, args.output, args.images_dir, sheet=args.sheet,
                     workers=args.workers, timeout=args.timeout, retries=args.retries,
                     dry_run=args.dry_run, resume=args.resume,
                     retry_transient_failures=args.retry_transient_failures)
    except (OSError, ValueError, KeyError) as exc:
        parser.error(str(exc))
    if args.dry_run:
        return 0
    return 2 if result["failed_unique_urls"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
