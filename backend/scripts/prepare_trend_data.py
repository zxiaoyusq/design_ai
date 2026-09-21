"""从文章信息表提取趋势描述、下载关联图片并输出 JSON / JSONL。

在项目根目录运行：python backend/scripts/prepare_trend_data.py
该命令只整理本地数据，不调用模型、不写入业务数据库。
"""

from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import UTC, date, datetime
import hashlib
from http.client import HTTPException
import json
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import time
from typing import Any, Callable
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qs, quote, urlsplit
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SCHEMA_VERSION = "trend_data_v1"
DESCRIPTION_FIELDS = (
    "title_zh", "summary_zh", "image_url", "image_width", "image_height",
    "primary_category", "subcategory", "tags", "confidence", "language_original",
    "release_time", "clust_status", "local_vl_info", "clustering_label",
)
# 新表新增的聚类标签原样输出；旧表可缺列，统一补 null，原有必填字段仍严格校验。
OPTIONAL_DESCRIPTION_FIELDS = frozenset({"clustering_label"})
IMAGE_EXTENSIONS = {
    "JPEG": ".jpg", "PNG": ".png", "WEBP": ".webp", "GIF": ".gif",
    "TIFF": ".tiff", "BMP": ".bmp", "AVIF": ".avif",
}


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _release_date(value: Any) -> str | None:
    """发布日期只保留源值的日历日期，不转换时区；空日期仍为 null。"""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return date.fromisoformat(str(value).strip()[:10]).isoformat()


def _number(value: Any, *, integer: bool = False) -> Any:
    """只转换无损数值；无法解释的原值继续保留，避免静默丢失描述。"""
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return value
    if not math.isfinite(number):
        return value
    if integer:
        return int(number) if number.is_integer() else value
    return number


def _labels(value: Any) -> Any:
    if value is None or value == "":
        return []
    if not isinstance(value, str):
        return value
    return [part.strip() for part in re.split(r"[,，]", value) if part.strip()]


def read_trends(
    source: Path, sheet_name: str | None = None, limit: int | None = None,
) -> tuple[list[dict[str, Any]], str]:
    """读取每行趋势并保留源行号；在任何下载或输出前校验字段和 ID。"""
    if limit is not None and limit < 1:
        raise ValueError("limit 必须大于 0")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            if len(workbook.sheetnames) != 1:
                raise ValueError("工作簿有多个工作表，请使用 --sheet 指定")
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"找不到工作表：{sheet_name}")
        rows = iter(workbook[sheet_name].iter_rows(values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        names = [name for name in headers if name]
        if len(names) != len(set(names)):
            raise ValueError("表头有重复字段")
        missing = set(("id", *DESCRIPTION_FIELDS)) - OPTIONAL_DESCRIPTION_FIELDS - set(headers)
        if missing:
            raise ValueError(f"缺少字段：{', '.join(sorted(missing))}")
        records: list[dict[str, Any]] = []
        seen_ids: set[str] = set()
        for row_number, values in enumerate(rows, start=2):
            if all(value is None or value == "" for value in values):
                continue
            row = dict(zip(headers, values))
            raw_id = row.get("id")
            if raw_id is None or not str(raw_id).strip():
                raise ValueError(f"第 {row_number} 行缺少 id")
            trend_id = str(raw_id).strip()
            if trend_id in seen_ids:
                raise ValueError(f"第 {row_number} 行 id 重复：{trend_id}")
            seen_ids.add(trend_id)
            record = {
                "id": trend_id,
                "source_row": row_number,
                **{field: _json_value(row.get(field)) for field in DESCRIPTION_FIELDS},
            }
            try:
                record["release_time"] = _release_date(row.get("release_time"))
            except ValueError as exc:
                raise ValueError(f"第 {row_number} 行 release_time 不是有效日期：{row.get('release_time')}") from exc
            for field in ("tags", "subcategory"):
                record[field] = _labels(record[field])
            for field in ("image_width", "image_height", "clust_status"):
                record[field] = _number(record[field], integer=True)
            record["confidence"] = _number(record["confidence"])
            vl_info = record["local_vl_info"]
            if isinstance(vl_info, str) and vl_info.strip():
                try:
                    parsed = json.loads(vl_info)
                    if not isinstance(parsed, (dict, list)):
                        raise ValueError("应为 JSON 对象或数组")
                    json.dumps(parsed, allow_nan=False)
                    record["local_vl_info"] = parsed
                except ValueError:
                    record["warnings"] = ["local_vl_info 不是有效 JSON 对象或数组，已保留原文"]
            # 重复链接仍占据各自位置，确保同一行图片数量、顺序和源表一一对应。
            urls = [url.strip() for url in str(record["image_url"] or "").split("||") if url.strip()]
            record["images"] = [
                {
                    "image_id": f"{trend_id}_{index:03d}", "trend_id": trend_id,
                    "index": index, "url": url, "local_path": None, "status": "pending",
                }
                for index, url in enumerate(urls, start=1)
            ]
            records.append(record)
        # 即使试跑前几行，也先验证整张表，避免完整运行时才发现重复 ID。
        return records[:limit] if limit is not None else records, sheet_name
    finally:
        workbook.close()


def _image_metadata(path: Path) -> dict[str, Any]:
    with Image.open(path) as image:
        image_format = image.format
        width, height = image.size
        image.verify()
    if image_format not in IMAGE_EXTENSIONS:
        raise ValueError(f"不支持的图片格式：{image_format}")
    # verify 对部分格式只检查文件头，load 再确认像素数据可解码。
    with Image.open(path) as image:
        image.load()
    with path.open("rb") as stream:
        digest = hashlib.file_digest(stream, "sha256").hexdigest()
    return {
        "format": image_format, "width": width, "height": height,
        "size_bytes": path.stat().st_size, "sha256": digest,
    }


def _image_stem(image: dict[str, Any], output_dir: Path, images_dir: Path | None = None) -> Path:
    """普通 ID 保持可读；其他 ID 加哈希避免路径穿越与清理后重名。"""
    trend_id = image["trend_id"]
    if re.fullmatch(r"[A-Za-z0-9_-]{1,100}", trend_id):
        directory = trend_id
    else:
        directory = "id_" + hashlib.sha256(trend_id.encode()).hexdigest()[:24]
    url_hash = hashlib.sha256(image["url"].encode()).hexdigest()[:20]
    return (images_dir or output_dir / "images") / directory / f"{image['index']:03d}_{url_hash}"


def _download_url(url: str) -> str:
    """该站的缩放接口易限流；仅解包已知媒体域名，直接下载参数明确指定的同一原图。"""
    parsed = urlsplit(url)
    if parsed.hostname == "materialdistrict.com" and parsed.path.rstrip("/") == "/_next/image":
        source = parse_qs(parsed.query).get("url", [""])[0]
        source_parts = urlsplit(source)
        if source_parts.scheme == "https" and source_parts.hostname == "media.materialdistrict.com":
            url = source
    # 解包后的原图路径可能含 ©、空格等字符；保留既有 % 转义和查询分隔符。
    return quote(url, safe=":/?#[]@!$&'()*+,;=%")


def download_image(
    image: dict[str, Any], output_dir: Path, timeout: float = 30,
    retries: int = 2, max_bytes: int = 50 * 1024 * 1024,
    opener: Callable[..., Any] | None = None,
    images_dir: Path | None = None,
) -> dict[str, Any]:
    """下载并校验原图；images_dir 可分离存储位置，返回路径仍相对 output_dir。

    可选 opener 仅替换连接方式，缓存和图片校验规则不变。
    """
    result = {**image, "status": "failed", "local_path": None}
    url = image["url"]
    try:
        request_url = _download_url(url)
    except ValueError as exc:
        return {**result, "error": str(exc)}
    image = {**image, "download_url": request_url}
    result["download_url"] = request_url
    stem = _image_stem(image, output_dir, images_dir)
    stem.parent.mkdir(parents=True, exist_ok=True)
    for extension in IMAGE_EXTENSIONS.values():
        cached = stem.with_suffix(extension)
        if not cached.is_file():
            continue
        try:
            metadata = _image_metadata(cached)
            if metadata["size_bytes"] > max_bytes:
                raise ValueError("缓存图片超出大小限制")
            return {
                **image, **metadata, "status": "downloaded", "cached": True,
                "download_url": None, "resolved_url": None,
                "local_path": Path(os.path.relpath(cached, output_dir)).as_posix(),
                "downloaded_at": datetime.fromtimestamp(cached.stat().st_mtime, UTC).isoformat(),
            }
        except (OSError, ValueError, UnidentifiedImageError, Image.DecompressionBombError):
            continue
    try:
        parsed = urlsplit(url)
        if parsed.scheme not in ("http", "https") or not parsed.hostname:
            raise ValueError("图片 URL 必须是有效的 HTTP(S) 地址")
    except ValueError as exc:
        return {**result, "error": str(exc)}
    for attempt in range(retries + 1):
        temporary_path: Path | None = None
        try:
            request = Request(request_url, headers={"User-Agent": "Mozilla/5.0", "Accept": "image/*,*/*;q=0.8"})
            with (opener or urlopen)(request, timeout=timeout) as response:
                resolved_url = response.url
                with tempfile.NamedTemporaryFile(dir=stem.parent, suffix=".part", delete=False) as stream:
                    temporary_path = Path(stream.name)
                    size = 0
                    while chunk := response.read(128 * 1024):
                        size += len(chunk)
                        if size > max_bytes:
                            raise ValueError(f"图片超过 {max_bytes} 字节限制")
                        stream.write(chunk)
            metadata = _image_metadata(temporary_path)
            destination = stem.with_suffix(IMAGE_EXTENSIONS[metadata["format"]])
            os.replace(temporary_path, destination)
            return {
                **image, **metadata, "status": "downloaded", "cached": False,
                "local_path": Path(os.path.relpath(destination, output_dir)).as_posix(),
                "resolved_url": resolved_url, "downloaded_at": datetime.now(UTC).isoformat(),
            }
        except (HTTPError, URLError, HTTPException, OSError, ValueError, UnidentifiedImageError, Image.DecompressionBombError) as exc:
            result["error"] = f"{type(exc).__name__}: {exc}"
            result["attempts"] = attempt + 1
            if isinstance(exc, HTTPError):
                result["http_status"] = exc.code
                exc.close()
                # 已明确不存在/拒绝访问的资源不会因立即重复请求而恢复。
                if exc.code not in (408, 429) and exc.code < 500:
                    break
            elif isinstance(exc, (ValueError, UnidentifiedImageError, Image.DecompressionBombError)):
                break
            if attempt < retries:
                time.sleep(min(2 ** attempt, 8))
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
    return result


def _atomic_text(path: Path, text: str) -> None:
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent, suffix=".tmp", delete=False) as stream:
            temporary_path = Path(stream.name)
            stream.write(text)
        os.replace(temporary_path, path)
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def write_outputs(output_dir: Path, records: list[dict[str, Any]], source: dict[str, Any]) -> dict[str, Any]:
    """JSON 是全量包，JSONL 每行一条趋势；报告单列失败资源方便补下载。"""
    images = [image for record in records for image in record["images"]]
    counts = {
        "trends": len(records), "image_references": len(images),
        "unique_image_urls": len({image["url"] for image in images}),
        "downloaded": sum(image["status"] == "downloaded" for image in images),
        "failed": sum(image["status"] == "failed" for image in images),
        "pending": sum(image["status"] == "pending" for image in images),
        "cached": sum(bool(image.get("cached")) for image in images),
        "trends_without_images": sum(not record["images"] for record in records),
        "trends_with_failed_images": sum(any(image["status"] == "failed" for image in record["images"]) for record in records),
    }
    metadata = {
        "schema_version": SCHEMA_VERSION, "generated_at": datetime.now(UTC).isoformat(),
        "source": source, "local_path_base": ".", "counts": counts,
    }
    _atomic_text(output_dir / "trends.json", json.dumps({**metadata, "trends": records}, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    _atomic_text(output_dir / "trends.jsonl", "".join(json.dumps(record, ensure_ascii=False, allow_nan=False) + "\n" for record in records))
    _atomic_text(output_dir / "download_report.json", json.dumps({
        **metadata,
        "failures": [image for image in images if image["status"] == "failed"],
        "warnings": [{"id": record["id"], "warnings": record["warnings"]} for record in records if record.get("warnings")],
    }, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    return counts


def main(argv: list[str] | None = None) -> int:
    """显式运行命令即开始本地整理；--dry-run 仅检查并展示处理范围。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=PROJECT_ROOT / "ref" / "文章信息表.xlsx")
    parser.add_argument("--output", type=Path, default=PROJECT_ROOT / "data" / "trend_data")
    parser.add_argument("--sheet", help="多工作表文件需显式指定工作表名称")
    parser.add_argument("--workers", type=int, default=12)
    parser.add_argument("--timeout", type=float, default=30)
    parser.add_argument("--retries", type=int, default=2, help="瞬时失败额外重试次数")
    parser.add_argument("--max-image-mb", type=int, default=50)
    parser.add_argument("--limit", type=int, help="仅整理前 N 条趋势，用于小批试跑")
    parser.add_argument("--dry-run", action="store_true", help="检查表格并展示范围，不写文件也不下载")
    args = parser.parse_args(argv)
    if args.workers < 1 or args.timeout <= 0 or args.retries < 0 or args.max_image_mb < 1:
        parser.error("workers、timeout、max-image-mb 必须大于 0，retries 必须不小于 0")
    input_path, output_dir = args.input.resolve(), args.output.resolve()
    records, sheet_name = read_trends(input_path, args.sheet, args.limit)
    images = [image for record in records for image in record["images"]]
    print(f"工作表 {sheet_name}：{len(records)} 条趋势，{len(images)} 个图片引用", flush=True)
    if args.dry_run:
        return 0
    output_dir.mkdir(parents=True, exist_ok=True)
    source_dir = output_dir / "source"
    source_dir.mkdir(exist_ok=True)
    # 保留源表副本及哈希，结构化转换后的值始终可回查原始单元格。
    source_copy = source_dir / input_path.name
    if input_path != source_copy:
        shutil.copy2(input_path, source_copy)
    source = {
        "file": source_copy.relative_to(output_dir).as_posix(), "sheet": sheet_name,
        "sha256": hashlib.sha256(source_copy.read_bytes()).hexdigest(),
        "description_fields": list(DESCRIPTION_FIELDS), "limit": args.limit,
    }
    readme = """# 趋势调研数据

- `trends.json`：含来源、统计信息和 `trends` 数组的完整数据包。
- `trends.jsonl`：每行一条完整趋势，适合逐条或分批输入 LLM。
- `images/<趋势ID>/<序号>_<URL哈希>.<实际格式>`：下载的原始图片。
- `download_report.json`：下载统计、失败 URL 与错误、字段解析提醒。
- `source/`：原始 Excel 副本；JSON 的 source 保留工作表名称和 SHA-256。

每条趋势保留原 id、Excel 行号 source_row 及全部 14 个指定描述字段。
clustering_label 保留源单元格原值，不拆分或重新聚类；旧表缺列或空单元格输出 null。
tags / subcategory 按中英文逗号转为数组，confidence 转为数值，clust_status
及可解析的源宽高转为整数，local_vl_info 解析为 JSON；解析失败保留原文。
空单元格保留 null，空标签为 []，release_time 只保留 YYYY-MM-DD，不转换时区。
image_url 保留原始多图字符串，images 按 || 分隔后的顺序记录每张图片。
源 image_width / image_height 不改写；每张实际图片尺寸见 images 的 width / height。

images 中 image_id / trend_id / index 标明关联，local_path 相对于本目录。
成功图片 status=downloaded，并含格式、宽高、字节数、SHA-256 和下载时间；
失败图片 status=failed、local_path=null，同时保留 URL、错误及已知 HTTP 状态码。
缓存复用时 cached=true，下载时间取缓存文件修改时间。
MaterialDistrict 缩放链接直接读取其 url 参数中明确指定的媒体原图，避免缩放接口限流；
url 保留表内链接，download_url 为请求地址，resolved_url 为本次请求的最终响应地址。
缓存复用不发起网络请求，download_url / resolved_url 均为 null，不伪造历史重定向信息。

向多模态模型输入时，需要读取 local_path 指向的文件并通过所用模型的图片接口附加；
仅在文本 Prompt 中放本地路径不会让远程模型读取该图片。按趋势分批输入，避免一次塞入全部数据。
整理过程不调用模型，源 local_vl_info 及 confidence 仅为表格已有数据，不代表重新验证的结论。

项目根目录执行 `conda run -n 314 python backend/scripts/prepare_trend_data.py` 可重新整理。
已完成且校验有效的图片会复用，失败或损坏图片重试。--dry-run 仅预览；--limit N 可试跑，
试跑建议同时指定独立 --output 目录，以免用子集覆盖全量索引。
运行过程中每完成 100 张保存一次索引。重新运行会重建索引，不会删除旧图片。
退出码 0 表示图片全部完成，2 表示有图片失败，1 表示输入或执行错误。
"""
    _atomic_text(output_dir / "README.md", readme)
    write_outputs(output_dir, records, source)
    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        jobs = {executor.submit(download_image, image, output_dir, args.timeout, args.retries, args.max_image_mb * 1024 * 1024): image for image in images}
        for completed, future in enumerate(as_completed(jobs), start=1):
            entry = jobs[future]
            entry.update(future.result())
            if completed % 100 == 0 or completed == len(images):
                counts = write_outputs(output_dir, records, source)
                print(f"图片 {completed}/{len(images)}：成功 {counts['downloaded']}，失败 {counts['failed']}，复用 {counts['cached']}", flush=True)
    counts = write_outputs(output_dir, records, source)
    print(f"已保存：{output_dir / 'trends.json'}", flush=True)
    return 2 if counts["failed"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
