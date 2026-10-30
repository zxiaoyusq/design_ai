"""PPT 原始证据、危险压缩包和完整原页渲染的小型离线验证。"""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import Mock, patch
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches
import pymupdf as fitz

from app.services.user_research import presentation as module


def make_pptx(path: Path) -> bytes:
    """两个虚构页面：同一原图在分组内、外各出现一次，保留两条出现关系。"""
    image = io.BytesIO()
    Image.new("RGB", (18, 12), (15, 100, 200)).save(image, format="PNG")
    blob = image.getvalue()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.3), Inches(6), Inches(0.5))
    title.text = "  虚构问题：喜欢哪一张？  \n回答：P1。"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "材质"
    table.cell(0, 1).text = "未提及"
    table.cell(1, 0).text = "颜色"
    table.cell(1, 1).text = "蓝色，\n不要灯线。"
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(3), Inches(1), Inches(1)).text = "P1"
    picture = group.shapes.add_picture(io.BytesIO(blob), Inches(3), Inches(3), Inches(2), Inches(1))
    picture.rotation = 15
    picture.crop_left = 0.1
    slide.notes_slide.notes_text_frame.text = "主持人提示：询问原因，尚无回答。"
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(blob), Inches(1), Inches(1), Inches(2), Inches(1))
    # 隐藏页依然属于原资料，必须进入渲染页数和后续检查。
    slide.element.set("show", "0")
    deck.save(path)
    return blob


def rewrite_zip(path: Path, transform) -> None:
    with zipfile.ZipFile(path) as source:
        entries = {entry.filename: source.read(entry) for entry in source.infolist()}
    transform(entries)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, payload in entries.items():
            target.writestr(name, payload)


def fake_pdf(source: Path, workspace: Path, count: int = 2) -> Path:
    path = workspace / "fixture.pdf"
    with fitz.open() as document:
        for _ in range(count):
            document.new_page(width=720, height=540)
        document.save(path)
    return path


class PresentationPreparationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / "fixture.pptx"
        self.blob = make_pptx(self.source)
        self.run = self.root / "run"

    def test_preserves_text_table_group_notes_and_original_media(self):
        with patch.object(module, "_convert_to_pdf", side_effect=fake_pdf):
            result = module.prepare_presentation(self.source, self.run)
        self.assertEqual(result["slideCount"], 2)
        self.assertEqual(result["sourceSha256"], hashlib.sha256(self.source.read_bytes()).hexdigest())
        page = result["slides"][0]
        self.assertEqual(page["texts"][0]["text"], "  虚构问题：喜欢哪一张？  \n回答：P1。")
        self.assertEqual([item["shapePath"] for item in page["texts"]], ["/0", "/1/table/0/0", "/1/table/0/1", "/1/table/1/0", "/1/table/1/1", "/2/0"])
        self.assertEqual(page["texts"][4]["text"], "蓝色，\n不要灯线。")
        self.assertIn("主持人提示：询问原因，尚无回答。", page["notes"])
        first, second = page["pictures"][0], result["slides"][1]["pictures"][0]
        self.assertEqual(first["shapePath"], "/2/1")
        self.assertEqual(first["rotation"], 15)
        self.assertEqual(first["crop"]["left"], 0.1)
        self.assertNotEqual(first["id"], second["id"])
        self.assertEqual(first["localPath"], second["localPath"])
        self.assertEqual(Path(first["localPath"]).read_bytes(), self.blob)
        self.assertEqual(first["sha256"], hashlib.sha256(self.blob).hexdigest())
        self.assertEqual(len(list((self.run / "media").iterdir())), 1)
        for slide in result["slides"]:
            self.assertTrue(Path(slide["renderPath"]).is_absolute())
            with Image.open(slide["renderPath"]) as image:
                self.assertEqual(max(image.size), 1600)
                self.assertEqual(image.format, "JPEG")
        self.assertEqual(json.loads((self.run / "inventory.json").read_text()), result)
        with patch.object(module, "_convert_to_pdf", side_effect=fake_pdf):
            repeated = module.prepare_presentation(self.source, self.root / "another-run")
        self.assertEqual(repeated["slides"][0]["pictures"][0]["id"], first["id"])

    def test_rejects_external_links_before_rendering(self):
        def external(entries):
            name = "ppt/slides/_rels/slide1.xml.rels"
            entries[name] = entries[name].replace(b"</Relationships>", b'<Relationship Id="rId999" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid/resource" TargetMode="External"/></Relationships>')
        rewrite_zip(self.source, external)
        with patch.object(module, "_convert_to_pdf") as render:
            with self.assertRaisesRegex(module.PresentationError, "外部链接"):
                module.prepare_presentation(self.source, self.run)
            render.assert_not_called()
        self.assertFalse((self.run / "inventory.json").exists())

    def test_rejects_corrupt_zip_path_traversal_macros_and_ole(self):
        good = self.source.read_bytes()
        for name in ("../outside", "ppt/vbaProject.bin", "ppt/embeddings/oleObject1.bin"):
            with self.subTest(name=name):
                self.source.write_bytes(good)
                rewrite_zip(self.source, lambda entries: entries.update({name: b"dummy"}))
                with patch.object(module, "_convert_to_pdf") as render:
                    with self.assertRaises(module.PresentationError):
                        module.prepare_presentation(self.source, self.run)
                    render.assert_not_called()
        self.source.write_bytes(b"not a zip")
        with self.assertRaisesRegex(module.PresentationError, "ZIP"):
            module.prepare_presentation(self.source, self.run)

    def test_rejects_embedded_svg_external_resource(self):
        rewrite_zip(self.source, lambda entries: entries.update({"ppt/media/unused.svg": b'<svg xmlns="http://www.w3.org/2000/svg"><image href="https://example.invalid/image.png"/></svg>'}))
        with self.assertRaisesRegex(module.PresentationError, "SVG.*外部资源"):
            module.prepare_presentation(self.source, self.run)

    def test_applies_size_and_page_limits_without_large_fixtures(self):
        for constant, limit in (("MAX_COMPRESSED_BYTES", 10), ("MAX_EXPANDED_BYTES", 10), ("MAX_MEDIA_BYTES", 10), ("MAX_SLIDES", 1)):
            with self.subTest(constant=constant), patch.object(module, constant, limit):
                with self.assertRaises(module.PresentationError):
                    module.prepare_presentation(self.source, self.run)

    def test_missing_renderer_and_page_mismatch_never_publish_success(self):
        with patch.object(module.shutil, "which", return_value=None):
            with self.assertRaisesRegex(module.PresentationError, "缺少 LibreOffice"):
                module.prepare_presentation(self.source, self.run)
        self.assertFalse((self.run / "inventory.json").exists())
        with patch.object(module, "_convert_to_pdf", side_effect=lambda source, workspace: fake_pdf(source, workspace, 1)):
            with self.assertRaisesRegex(module.PresentationError, "页数不一致"):
                module.prepare_presentation(self.source, self.run)
        self.assertFalse((self.run / "inventory.json").exists())
        self.assertFalse((self.run / "media").exists())

    def test_renderer_timeout_only_kills_its_own_process_group(self):
        process = Mock(pid=12345)
        process.wait.side_effect = [subprocess.TimeoutExpired("soffice", 180), -9]
        workspace = self.root / "renderer"
        workspace.mkdir()
        with patch.object(module.shutil, "which", return_value="/usr/bin/soffice"), patch.object(module.subprocess, "Popen", return_value=process) as popen, patch.object(module.os, "killpg") as kill:
            with self.assertRaisesRegex(module.PresentationError, "180 秒"):
                module._convert_to_pdf(self.source, workspace)
        kill.assert_called_once_with(12345, module.signal.SIGKILL)
        self.assertTrue(popen.call_args.kwargs["start_new_session"])
        self.assertIn(f"-env:UserInstallation={(workspace / 'office-profile').as_uri()}", popen.call_args.args[0])

    @unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "未安装 LibreOffice，离线错误路径仍已验证")
    def test_real_libreoffice_renders_two_synthetic_pages_including_hidden(self):
        result = module.prepare_presentation(self.source, self.run)
        self.assertEqual(len(result["slides"]), 2)
        for page in result["slides"]:
            self.assertGreater(Path(page["renderPath"]).stat().st_size, 1000)


if __name__ == "__main__":
    unittest.main()
