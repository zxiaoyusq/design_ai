"""受限读取 PPTX 原始证据，并通过独立 LibreOffice 进程渲染原页。"""
from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path, PurePosixPath
import posixpath
import re
import shutil
import signal
import stat
import subprocess
import tempfile
from urllib.parse import unquote, urlsplit
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import pymupdf as fitz


MAX_COMPRESSED_BYTES = 100 * 1024 * 1024
MAX_EXPANDED_BYTES = 1024 * 1024 * 1024
MAX_MEDIA_BYTES = 50 * 1024 * 1024
MAX_XML_BYTES = 20 * 1024 * 1024
MAX_SLIDES = 80
MAX_ZIP_ENTRIES = 20_000
RENDER_TIMEOUT_SECONDS = 180
RENDER_LONG_EDGE = 1600
_PML = "http://schemas.openxmlformats.org/presentationml/2006/main"
_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_UNSAFE_CONTENT = ("macroenabled", "vbaproject", "oleobject", "activex")


class PresentationError(ValueError):
    """PPTX 不安全、不受支持或无法完整渲染；调用方不得将其标为已完成。"""


def _read_member(archive: zipfile.ZipFile, member: zipfile.ZipInfo, limit: int) -> bytes:
    if member.file_size > limit:
        raise PresentationError(f"PPTX 内部文件超出大小限制：{member.filename}")
    with archive.open(member) as stream:
        data = stream.read(limit + 1)
    if len(data) > limit or len(data) != member.file_size:
        raise PresentationError(f"PPTX 内部文件大小异常：{member.filename}")
    return data


def _xml(data: bytes, name: str):
    # 禁用实体和网络，再拒绝 DTD，避免后续 Office 解析器解释外部实体。
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    try:
        root = etree.fromstring(data, parser)
    except etree.XMLSyntaxError as exc:
        raise PresentationError(f"PPTX XML 无效：{name}") from exc
    if root.getroottree().docinfo.doctype:
        raise PresentationError(f"PPTX 不允许 XML DTD：{name}")
    return root


def _validate_relationships(root, name: str, names: set[str]) -> None:
    source = name.removesuffix(".rels").replace("_rels/", "")
    base = posixpath.dirname(source)
    for relation in root:
        if relation.tag != f"{{{_REL}}}Relationship":
            raise PresentationError(f"PPTX 关系文件格式无效：{name}")
        mode = relation.get("TargetMode", "Internal").lower()
        target = unquote(relation.get("Target", ""))
        relation_type = relation.get("Type", "").lower()
        if mode != "internal" or urlsplit(target).scheme or target.startswith("//") or "\\" in target:
            raise PresentationError(f"PPTX 含外部链接，不能离线渲染：{name}")
        if any(value in relation_type for value in _UNSAFE_CONTENT) or relation_type.endswith("/package"):
            raise PresentationError(f"PPTX 含宏或嵌入对象：{name}")
        target_path = target.split("#", 1)[0]
        resolved = posixpath.normpath(target_path.lstrip("/") if target_path.startswith("/") else posixpath.join(base, target_path))
        if resolved.startswith("../") or resolved == ".." or resolved not in names:
            raise PresentationError(f"PPTX 内部链接无效：{name}")


def _validate_svg(root, name: str) -> None:
    # SVG 可包含无需 OPC relationship 的资源引用，不能只检查 .rels。
    for node in root.iter():
        if not isinstance(node.tag, str):
            continue
        if etree.QName(node).localname.lower() in {"script", "foreignobject"}:
            raise PresentationError(f"PPTX SVG 含活动内容：{name}")
        for key, value in node.attrib.items():
            if etree.QName(key).localname.lower() in {"href", "src", "base"} and value and not value.startswith("#"):
                raise PresentationError(f"PPTX SVG 含外部资源：{name}")
        for value in [*node.attrib.values(), node.text or ""]:
            if "@import" in value.lower():
                raise PresentationError(f"PPTX SVG 含外部样式：{name}")
            for target in re.findall(r"url\s*\((.*?)\)", value, flags=re.IGNORECASE | re.DOTALL):
                if not target.strip().strip("\"'").startswith("#"):
                    raise PresentationError(f"PPTX SVG 含外部资源：{name}")


def _validate_archive(data: bytes) -> None:
    """全包限额并逐项验证；不 extractall，也不信任上传方提供的 ZIP 路径。"""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_ENTRIES:
                raise PresentationError("PPTX 内部文件数量超限")
            names: set[str] = set()
            expanded = 0
            for member in members:
                name = member.filename
                segments = PurePosixPath(name).parts
                if (not name or name.startswith("/") or "\\" in name or ":" in name
                        or any(part in {".", ".."} for part in name.split("/"))
                        or any(ord(char) < 32 for char in name)
                        or stat.S_ISLNK(member.external_attr >> 16)):
                    raise PresentationError(f"PPTX 含非法 ZIP 路径：{name}")
                if name in names:
                    raise PresentationError(f"PPTX 含重复 ZIP 文件：{name}")
                names.add(name)
                if member.flag_bits & 1:
                    raise PresentationError("PPTX 不支持加密 ZIP 文件")
                lowered = name.lower()
                if any(value in lowered for value in _UNSAFE_CONTENT) or "embeddings" in segments:
                    raise PresentationError(f"PPTX 含宏或嵌入对象：{name}")
                expanded += member.file_size
                if expanded > MAX_EXPANDED_BYTES:
                    raise PresentationError("PPTX 展开大小超过 1 GB 限制")
                if lowered.startswith("ppt/media/") and member.file_size > MAX_MEDIA_BYTES:
                    raise PresentationError("PPTX 单个媒体超过 50 MB 限制")
            if not {"[Content_Types].xml", "ppt/presentation.xml", "_rels/.rels"} <= names:
                raise PresentationError("上传文件不是有效的 PPTX")
            slide_count = None
            for member in members:
                if member.is_dir():
                    continue
                name = member.filename
                is_xml = name.lower().endswith((".xml", ".rels", ".svg"))
                payload = _read_member(archive, member, MAX_XML_BYTES if is_xml else MAX_EXPANDED_BYTES)
                if not is_xml:
                    continue
                root = _xml(payload, name)
                if name == "[Content_Types].xml":
                    for node in root:
                        content_type = node.get("ContentType", "").lower()
                        if any(value in content_type for value in _UNSAFE_CONTENT):
                            raise PresentationError("PPTX 含宏或嵌入对象内容类型")
                elif name == "ppt/presentation.xml":
                    slide_count = len(root.findall(f"{{{_PML}}}sldIdLst/{{{_PML}}}sldId"))
                if name.lower().endswith(".rels"):
                    _validate_relationships(root, name, names)
                if name.lower().endswith(".svg"):
                    _validate_svg(root, name)
                if any(isinstance(node.tag, str) and etree.QName(node).localname.lower() in {"oleobj", "control"} for node in root.iter()):
                    raise PresentationError(f"PPTX 含 OLE 或活动控件：{name}")
            if slide_count is None or not 1 <= slide_count <= MAX_SLIDES:
                raise PresentationError("PPTX 必须包含 1 至 80 页")
    except (zipfile.BadZipFile, RuntimeError, NotImplementedError, OSError, EOFError) as exc:
        raise PresentationError("PPTX ZIP 文件损坏或使用不支持的压缩方式") from exc


def _bounds(shape) -> dict:
    # 与原 PPT inspector 一致：EMU 单位；组内形状保留父组坐标及递归路径。
    return {key: int(getattr(shape, key)) for key in ("left", "top", "width", "height")}


def _collect_shape(shape, shape_path: str, page: int, media: dict, texts: list, pictures: list) -> None:
    if shape.has_text_frame and shape.text != "":
        texts.append({"shapePath": shape_path, "text": shape.text, **_bounds(shape)})
    if shape.has_table:
        table = shape.table
        widths = [int(column.width) for column in table.columns]
        heights = [int(row.height) for row in table.rows]
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                if cell.is_spanned or cell.text == "":
                    continue
                texts.append({
                    "shapePath": f"{shape_path}/table/{row_index}/{column_index}", "text": cell.text,
                    "left": int(shape.left) + sum(widths[:column_index]),
                    "top": int(shape.top) + sum(heights[:row_index]),
                    "width": sum(widths[column_index:column_index + cell.span_width]),
                    "height": sum(heights[row_index:row_index + cell.span_height]),
                })
    if shape.element.tag == f"{{{_PML}}}pic":
        relationship_id = shape.element.blip_rId
        if not relationship_id:
            raise PresentationError(f"第 {page} 页图片 {shape_path} 缺少内嵌原图")
        image_part = shape.part.related_part(relationship_id)
        source = media.get(str(image_part.partname).lstrip("/"))
        if source is None:
            raise PresentationError(f"第 {page} 页图片 {shape_path} 未指向原始媒体")
        identifier = hashlib.sha256(f"{page}:{shape_path}:{source['sha256']}".encode()).hexdigest()
        pictures.append({
            "id": "ppt_" + identifier[:32], "shapePath": shape_path,
            "sha256": source["sha256"], "localPath": source["localPath"],
            **_bounds(shape), "rotation": float(shape.rotation),
            "crop": {side: float(getattr(shape, f"crop_{side}")) for side in ("left", "top", "right", "bottom")},
        })
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for index, child in enumerate(shape.shapes):
            _collect_shape(child, f"{shape_path}/{index}", page, media, texts, pictures)


def _extract(data: bytes, stage: Path, run: Path) -> list[dict]:
    media = {}
    (stage / "media").mkdir()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for member in archive.infolist():
            if member.filename.lower().startswith("ppt/media/") and not member.is_dir():
                blob = _read_member(archive, member, MAX_MEDIA_BYTES)
                digest = hashlib.sha256(blob).hexdigest()
                suffix = PurePosixPath(member.filename).suffix.lower()
                if not re.fullmatch(r"\.[a-z0-9]{1,10}", suffix):
                    suffix = ".bin"
                filename = digest + suffix
                (stage / "media" / filename).write_bytes(blob)
                media[member.filename] = {"sha256": digest, "localPath": str(run / "media" / filename)}
    try:
        presentation = Presentation(io.BytesIO(data))
        slides = []
        for number, slide in enumerate(presentation.slides, 1):
            texts, pictures, notes = [], [], []
            for index, shape in enumerate(slide.shapes):
                _collect_shape(shape, f"/{index}", number, media, texts, pictures)
            # 不创建不存在的 notes part；备注仅是原始资料，角色留给后续语义复核。
            if slide.has_notes_slide:
                for shape in slide.notes_slide.shapes:
                    if shape.has_text_frame and shape.text != "":
                        notes.append(shape.text)
            slides.append({"number": number, "texts": texts, "notes": notes, "pictures": pictures})
        return slides
    except PresentationError:
        raise
    except (ValueError, KeyError, AttributeError, TypeError, IndexError) as exc:
        raise PresentationError("PPTX 页面结构无法完整读取") from exc


def _convert_to_pdf(source: Path, workspace: Path) -> Path:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise PresentationError("缺少 LibreOffice Impress 渲染器，无法核对原页")
    profile = workspace / "office-profile"
    (profile / "user").mkdir(parents=True)
    # 系统 LibreOffice 7.3 尚不支持命令行 JSON filter options，设置本次专用
    # profile 同时兼容旧版和新版；不修改上传文件来取消隐藏属性。
    (profile / "user" / "registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Common/Filter/PDF/Export">'
        '<prop oor:name="ExportHiddenSlides" oor:op="fuse"><value>true</value></prop>'
        '</item></oor:items>', encoding="utf-8")
    output = workspace / "pdf"
    output.mkdir()
    # 独立 profile 防止连接到服务器已有 Office 实例；隐藏页也必须输出。
    command = [executable, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--invisible",
               "--nologo", "--nodefault", "--norestore", "--convert-to", "pdf:impress_pdf_Export",
               "--outdir", str(output), str(source)]
    try:
        with (workspace / "render.log").open("wb") as log:
            process = subprocess.Popen(command, cwd=workspace, stdin=subprocess.DEVNULL,
                                       stdout=log, stderr=subprocess.STDOUT, start_new_session=True)
            try:
                returncode = process.wait(timeout=RENDER_TIMEOUT_SECONDS)
            except subprocess.TimeoutExpired as exc:
                # 只终止本次新建的进程组，绝不 killall 或清理共享 Office profile。
                try:
                    os.killpg(process.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass
                process.wait()
                raise PresentationError("PPTX 原页渲染超过 180 秒") from exc
    except OSError as exc:
        raise PresentationError("LibreOffice Impress 渲染器无法启动") from exc
    pdf = output / (source.stem + ".pdf")
    if returncode != 0 or not pdf.is_file() or pdf.stat().st_size == 0:
        raise PresentationError("LibreOffice 未生成有效 PDF，原页渲染失败")
    return pdf


def _render_pages(pdf: Path, output: Path, expected_pages: int) -> list[Path]:
    try:
        with fitz.open(pdf) as document:
            if document.page_count != expected_pages:
                raise PresentationError(f"渲染页数不一致：PPTX {expected_pages} 页，PDF {document.page_count} 页")
            output.mkdir()
            paths = []
            for index, page in enumerate(document, 1):
                longest = max(page.rect.width, page.rect.height)
                if longest <= 0:
                    raise PresentationError(f"第 {index} 页 PDF 尺寸无效")
                scale = RENDER_LONG_EDGE / longest
                pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False)
                path = output / f"slide-{index:03d}.jpg"
                pixmap.save(path, jpg_quality=90)
                paths.append(path)
            return paths
    except PresentationError:
        raise
    except (RuntimeError, ValueError, OSError) as exc:
        raise PresentationError("PDF 无法完整渲染为原页图片") from exc


def prepare_presentation(path: Path, run: Path) -> dict:
    """冻结本次 PPTX 字节、保存原媒体并渲染全部原页；成功才发布 inventory。

    调用方提供独立任务目录。返回绝对媒体/原页路径；文本保留原文和顺序，
    shapePath 使用 inspector 的 /0/1 递归索引，表格追加 /table/行/列。
    坐标为 PPT 原生 EMU（组内为父组坐标），不是截图像素；不执行语义判断。
    """
    path, run = Path(path), Path(run).resolve()
    if path.suffix.lower() != ".pptx":
        raise PresentationError("仅支持 .pptx 演示文稿")
    if not path.is_file():
        raise PresentationError("PPTX 源文件不存在")
    if path.stat().st_size > MAX_COMPRESSED_BYTES:
        raise PresentationError("PPTX 压缩文件超过 100 MB 限制")
    with path.open("rb") as stream:
        data = stream.read(MAX_COMPRESSED_BYTES + 1)
    if len(data) > MAX_COMPRESSED_BYTES:
        raise PresentationError("PPTX 压缩文件超过 100 MB 限制")
    _validate_archive(data)
    run.mkdir(parents=True, exist_ok=True)
    if any(os.path.lexists(run / name) for name in ("media", "page-renders", "inventory.json")):
        raise PresentationError("任务目录已有 PPT 准备产物，请使用新的运行目录")
    with tempfile.TemporaryDirectory(prefix=".presentation-", dir=run) as temporary:
        stage = Path(temporary)
        slides = _extract(data, stage, run)
        source = stage / "source.pptx"
        source.write_bytes(data)
        pdf = _convert_to_pdf(source, stage)
        renders = _render_pages(pdf, stage / "page-renders", len(slides))
        for slide, rendered in zip(slides, renders, strict=True):
            slide["renderPath"] = str(run / "page-renders" / rendered.name)
        result = {"sourceSha256": hashlib.sha256(data).hexdigest(), "slideCount": len(slides), "slides": slides}
        (stage / "inventory.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        for name in ("media", "page-renders", "inventory.json"):
            (stage / name).rename(run / name)
        return result
