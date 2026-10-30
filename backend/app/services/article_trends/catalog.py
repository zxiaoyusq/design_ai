"""从源 Excel 读取文字，单独关联已有图片索引；模型输入不包含图片。"""

from __future__ import annotations

from datetime import date, datetime
import json
from pathlib import Path
from threading import RLock
from urllib.parse import urlsplit, quote

from openpyxl import load_workbook

from app.schemas.article_trends import ArticleFilters, ClusterRequest

PROJECT = Path(__file__).resolve().parents[4]


def safe_url(value):
    text = str(value or "").strip()
    try:
        return text if urlsplit(text).scheme in ("http", "https") and urlsplit(text).hostname else None
    except ValueError:
        return None


class ArticleCatalog:
    """按源文件版本缓存完整记录；分页响应只带摘要和封面，详情再展开图片。"""

    def __init__(self, root=PROJECT):
        self.root = Path(root)
        self.excel = self.root / "ref/趋势信息0930.xlsx"
        self.image_index = self.root / "data/trend_data/trend_info_0930/trends.json"
        self.images_root = self.root / "data/trend_data/images"
        self._stamp = None
        self._records = []
        self._images = {}
        self._lock = RLock()

    def records(self):
        with self._lock:
            paths = (self.excel, self.image_index)
            stamp = tuple((p.stat().st_mtime_ns, p.stat().st_size) if p.exists() else None for p in paths)
            if stamp != self._stamp:
                self._load()
                self._stamp = stamp
            return self._records

    def _load(self):
        workbook = load_workbook(self.excel, read_only=True, data_only=True)
        records = []
        seen = set()
        try:
            sheet = workbook["导出结果"] if "导出结果" in workbook.sheetnames else workbook.active
            # 导出源表的 dimension=A1 不代表真实内容范围。
            sheet.reset_dimensions()
            rows = iter(sheet.iter_rows(values_only=True))
            headers = [str(v or "").strip() for v in next(rows)]
            for number, values in enumerate(rows, 2):
                raw = dict(zip(headers, values))
                if raw.get("id") is None:
                    continue
                aid = str(raw["id"]).strip()
                if aid in seen:
                    raise ValueError("源表文章 ID 重复：" + aid)
                seen.add(aid)
                title = str(raw.get("title_zh") or raw.get("title") or "")
                summary = str(raw.get("summary_zh") or "")
                dt = raw.get("release_time")
                if isinstance(dt, (date, datetime)):
                    published = dt.isoformat()[:10]
                else:
                    try:
                        published = date.fromisoformat(str(dt)[:10]).isoformat()
                    except ValueError:
                        published = None
                detail = safe_url(raw.get("detail_url"))
                site = safe_url(raw.get("source_url"))
                source = urlsplit(detail or site or "").hostname or "来源未提供"
                records.append({
                    "article_id": aid, "title": title, "summary": summary,
                    "published_at": published, "source": source,
                    "source_url": detail or site, "detail_url": detail,
                    "category": str(raw.get("primary_category") or "未分类"),
                    "text_kind": "source_paraphrase", "source_row": number,
                })
        finally:
            workbook.close()
        images = {}
        if self.image_index.is_file():
            data = json.loads(self.image_index.read_text(encoding="utf-8"))
            images = {str(a["id"]): a.get("images", []) for a in data["trends"]}
        # 同日保留 Excel 行序，便于“最新 N 篇”的稳定冻结。
        records.sort(key=lambda a: a["published_at"] or "", reverse=True)
        self._records, self._images = records, images

    def filtered(self, filters: ArticleFilters):
        query = filters.q.strip().casefold()
        start = filters.date_from.isoformat() if filters.date_from else None
        end = filters.date_to.isoformat() if filters.date_to else None
        return [a for a in self.records() if
                (not query or query in (a["title"] + "\n" + a["summary"]).casefold())
                and (not filters.source or a["source"] == filters.source)
                and (not filters.category or a["category"] == filters.category)
                and (not start or (a["published_at"] and a["published_at"] >= start))
                and (not end or (a["published_at"] and a["published_at"] <= end))]

    def selected(self, request: ClusterRequest):
        if request.article_ids is not None:
            ids = set(request.article_ids)
            if not ids:
                raise ValueError("请至少勾选一篇文章")
            records = [a for a in self.records() if a["article_id"] in ids]
            if len(records) != len(ids):
                raise ValueError("勾选范围包含不存在的文章，请刷新文章库")
            # 手动范围忽略后续关键词/分类变化，但趋势日期仍是本次输入的硬边界。
            dated_ids = {a["article_id"] for a in self.filtered(ArticleFilters(
                date_from=request.filters.date_from, date_to=request.filters.date_to))}
            records = [a for a in records if a["article_id"] in dated_ids]
        else:
            records = self.filtered(request.filters)
        return records[:request.limit] if request.limit else records

    def semantic_records(self, selected):
        # 白名单保证分类、标签、图片 URL、图片 DNA 永不进入语义输入。
        return [{k: a[k] for k in ("article_id", "title", "summary", "text_kind", "published_at")}
                | {"source_url": a.get("detail_url")} for a in selected]

    def public(self, article, detail=False):
        available = self.available_images(article["article_id"], limit=None if detail else 1)
        refs = self._images.get(article["article_id"], [])
        return {k: article.get(k) for k in (
            "article_id", "title", "summary", "published_at", "source", "source_url", "category")
        } | {"image_count": len(refs), "images": [
            {key: image[key] for key in ("image_id", "url")} for image in available
        ]}

    def available_images(self, aid, limit=None):
        """复用图片路由的本地安全校验；原链接仅供跨文章去重，不作为展示地址。"""
        self.records()
        images = []
        article_id = quote(aid, safe="")
        for index, ref in enumerate(self._images.get(aid, [])):
            try:
                self.image(aid, index)
            except FileNotFoundError:
                continue
            images.append({
                "image_id": str(ref.get("image_id", index)),
                "url": f"/api/v1/article-trends/articles/{article_id}/images/{index}",
                "source_url": str(ref.get("url") or "").strip(),
                "source_sha256": str(ref.get("sha256") or "").strip(),
            })
            if limit is not None and len(images) >= limit:
                break
        return images

    def listing(self, filters, page=1, page_size=30):
        records = self.filtered(filters)
        return {"total": len(records), "page": page, "page_size": page_size,
                "items": [self.public(a) for a in records[(page - 1)*page_size:page*page_size]]}

    def article(self, aid):
        item = next((a for a in self.records() if a["article_id"] == aid), None)
        if item is None:
            raise FileNotFoundError("文章不存在")
        return self.public(item, detail=True)

    def image(self, aid, index):
        self.records()
        refs = self._images.get(aid, [])
        if index < 0 or index >= len(refs) or refs[index].get("status") != "downloaded":
            raise FileNotFoundError("图片不存在")
        relative = refs[index].get("local_path")
        if not relative:
            raise FileNotFoundError("图片不存在")
        path = (self.image_index.parent / relative).resolve()
        if not path.is_relative_to(self.images_root.resolve()) or not path.is_file():
            raise FileNotFoundError("图片不可用")
        return path

    def facets(self, version):
        rows = self.records()
        dates = [a["published_at"] for a in rows if a["published_at"]]
        return {"total": len(rows), "sources": sorted({a["source"] for a in rows}),
                "categories": sorted({a["category"] for a in rows}),
                "min_date": min(dates, default=None), "max_date": max(dates, default=None), "skill_version": version}
