"""微型合成表验证文本隔离、来源覆盖与归并边界，不调用模型。"""

import copy
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts/pipeline.py"
spec = importlib.util.spec_from_file_location("trend_text_pipeline", SCRIPT)
p = importlib.util.module_from_spec(spec)
spec.loader.exec_module(p)


def workbook(path, rows=None, wrong_dimensions=False):
    book = Workbook()
    sheet = book.active
    sheet.append(["id", "title", "title_zh", "summary_zh", "source_url", "detail_url", "local_vl_info", "image_url", "tags"])
    for row in rows or [
        ["1", "A", "灯具", "主体保持哑光，边缘形成亮面细节。", "https://site.example", "https://site.example/a", "视觉秘密：大红色", "https://images.invalid/a.jpg", "标签秘密：高端"],
        ["2", "B", "座椅", "低光泽椅面与明亮细节形成对比。", "https://site.example", "https://site.example/b", "视觉秘密：蓝色", "https://images.invalid/b.jpg", "标签秘密：热门"],
        ["3", "C", "电池新闻", "电池容量提升20%，未描述外观。", "https://site.example", "https://site.example/c", "视觉秘密：黄色", "https://images.invalid/c.jpg", "标签秘密：科技"],
    ]:
        sheet.append(row)
    book.save(path)
    if wrong_dimensions:
        with ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b'ref="A1:I4"', b'ref="A1"')
        with ZipFile(path, "w") as archive:
            for name, content in entries.items():
                archive.writestr(name, content)


def theme(cid):
    return {"cluster_id": cid, "title": "低光泽主体与亮面细节", "definition": "主体和细节的光泽角色不同。",
            "design_question": "如何以局部光泽对比形成重点？", "dimensions": ["surface", "composition"],
            "boundaries": {"shared_design_mechanism": "低反射主体与亮点的主次关系", "includes": ["主体与细节反差"], "excludes": ["整面镜面"]},
            "hypotheses": [], "unknowns": ["真实反射率未知"]}


ANALYSIS = {"actor": "synthetic_fixture", "model_id": None, "model_version": None, "completed_at": "2026-09-30T12:00:00+00:00"}


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / "source.xlsx"
        workbook(self.source)
        self.run = self.root / "run"
        p.prepare(self.source, self.run)
        self.request = p.read(self.run / "batches/B0001/request.json")
        self.catalogue = p.read(self.run / "catalogue.json")

    def tearDown(self):
        self.tmp.cleanup()

    def proposal(self):
        clusters = []
        for index in (1, 2):
            cluster = theme(f"B0001-C{index}")
            eid = f"T:{index}:summary_zh"
            text = next(item["text"] for item in self.catalogue["evidence"] if item["evidence_id"] == eid)
            cluster["members"] = [{"point_id": f"B0001-P{index}", "subject": "主体与细节", "feature": "主次光泽反差",
                                   "dimension": "surface", "condition": None,
                                   "evidence_id": eid, "quote": text, "reason": "表内明确描述主次光泽关系"}]
            clusters.append(cluster)
        return {"batch_id": "B0001", "request_sha256": self.request["request_sha256"], "analysis": ANALYSIS,
                "clusters": clusters, "coverage": [{"source_id": f"T:{index}", "status": "clustered" if index < 3 else "non_aesthetic",
                  "reason": "支持局部光泽对比" if index < 3 else "仅有容量数据，没有审美依据", "residuals": []} for index in (1, 2, 3)]}

    def accept_and_collect(self):
        path = self.root / "proposal.json"
        p.write(path, self.proposal())
        p.check_batch(self.run, "B0001", path)
        p.collect(self.run)
        cluster = theme("TC01")
        cluster["candidate_ids"] = ["B0001-C1", "B0001-C2"]
        plan = {"collection_sha256": p.digest(self.run / "collection.json"), "analysis": ANALYSIS, "clusters": [cluster]}
        path = self.root / "plan.json"
        p.write(path, plan)
        return path

    def test_wrong_dimension_and_text_isolation(self):
        path = self.root / "wrong.xlsx"
        workbook(path, wrong_dimensions=True)
        catalogue, links = p.inventory(path)
        self.assertEqual(len(catalogue["articles"]), 3)
        serialized = json.dumps(self.request, ensure_ascii=False)
        for marker in ("视觉秘密", "标签秘密", "images.invalid", "site.example"):
            self.assertNotIn(marker, serialized)
        self.assertNotIn("视觉秘密", json.dumps(catalogue, ensure_ascii=False))
        self.assertIn("images.invalid", json.dumps(links))

    def test_same_publisher_not_same_article_and_duplicate_summary_collapses(self):
        self.assertEqual(len({a["source_group_id"] for a in self.catalogue["articles"]}), 3)
        path = self.root / "duplicates.xlsx"
        workbook(path, [["a", "A", "标题甲", "重复摘要", "https://a.example", "https://a.example/one", "", "", ""],
                        ["b", "B", "标题乙", "不同摘要", "https://a.example", "https://a.example/one?utm_source=x", "", "", ""],
                        ["c", "C", "标题丙", "不同摘要", "https://b.example", "https://b.example/other", "", "", ""]])
        catalogue, _ = p.inventory(path)
        self.assertEqual(len({a["source_group_id"] for a in catalogue["articles"]}), 1)

    def test_hash_routed_stories_stay_distinct_and_detail_domain_fallback(self):
        path = self.root / "stories.xlsx"
        workbook(path, [["a", "A", "标题甲", "摘要甲", None, "https://news.example/feed#story-one", "", "", ""],
                        ["b", "B", "标题乙", "摘要乙", None, "https://news.example/feed#story-two", "", "", ""]])
        catalogue, _ = p.inventory(path)
        self.assertEqual(len({a["source_group_id"] for a in catalogue["articles"]}), 2)
        self.assertEqual({a["publisher_group"] for a in catalogue["articles"]}, {"news.example"})

    def test_lost_source_wrong_status_or_invented_quote_rejected(self):
        for mutation in ("lost", "status", "quote", "foreign"):
            proposal = self.proposal()
            if mutation == "lost":
                proposal["coverage"].pop()
            elif mutation == "status":
                proposal["coverage"][2]["status"] = "clustered"
            elif mutation == "quote":
                proposal["clusters"][0]["members"][0]["quote"] = "模型虚构描述"
            else:
                proposal["clusters"][0]["members"][0]["evidence_id"] = "T:99:summary_zh"
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                p.validate_batch(proposal, self.request, self.catalogue)

    def test_end_to_end_keeps_evidence_and_exclusions(self):
        plan = self.accept_and_collect()
        out = self.root / "final"
        p.finalize(self.run, plan, out)
        result = p.read(out / "trend-clusters.json")
        self.assertEqual(result["schema_version"], "prism.trends.semantic_clusters.v0.1")
        self.assertEqual(result["clusters"][0]["source_ids"], ["T:1", "T:2"])
        self.assertEqual(result["clusters"][0]["source_group_count"], 2)
        self.assertEqual(len(result["clusters"][0]["lineage"]), 2)
        self.assertFalse(result["visual_analysis_performed"])
        self.assertEqual(p.read(out / "coverage.json")[2]["status"], "non_aesthetic")
        self.assertEqual(p.read(out / "execution.json")["trace_gaps"], ["B0001.model_id", "B0001.model_version", "merge.model_id", "merge.model_version"])
        with self.assertRaises(ValueError):
            p.finalize(self.run, plan, out)

    def test_plan_missing_or_repeated_candidate_rejected(self):
        plan_path = self.accept_and_collect()
        plan = p.read(plan_path)
        for ids in (["B0001-C1"], ["B0001-C1", "B0001-C1"]):
            changed = copy.deepcopy(plan)
            changed["clusters"][0]["candidate_ids"] = ids
            p.write(plan_path, changed)
            with self.assertRaises(ValueError):
                p.finalize(self.run, plan_path, self.root / "invalid")
            self.assertFalse((self.root / "invalid").exists())

    def test_revision_invalidates_collection_and_keeps_history(self):
        plan_path = self.accept_and_collect()
        proposal = self.proposal()
        proposal["clusters"][0]["unknowns"].append("样本实现范围未验证")
        p.write(self.root / "revised.json", proposal)
        p.check_batch(self.run, "B0001", self.root / "revised.json")
        self.assertEqual(len(list((self.run / "batches/B0001/history").glob("*.json"))), 2)
        with self.assertRaisesRegex(ValueError, "变化"):
            p.finalize(self.run, plan_path, self.root / "invalid")
        p.collect(self.run)
        with self.assertRaisesRegex(ValueError, "当前候选集合"):
            p.finalize(self.run, plan_path, self.root / "invalid")

    def test_frozen_input_tampering_rejected(self):
        altered = copy.deepcopy(self.catalogue)
        altered["evidence"][0]["text"] = "changed"
        p.write(self.run / "catalogue.json", altered)
        with self.assertRaisesRegex(ValueError, "冻结输入"):
            p.load_run(self.run)

    def test_budgets_no_truncation_and_no_partial_directory(self):
        items = [{"source_id": str(i), "text": "a" * 80} for i in range(3)]
        groups = p.pack(items, "prompt", "articles", 260)
        self.assertGreater(len(groups), 1)
        self.assertEqual([item for group in groups for item in group], items)
        with self.assertRaises(ValueError):
            p.prepare(self.source, self.root / "too-small", max_chars=20)
        self.assertFalse((self.root / "too-small").exists())

    def test_empty_text_has_explicit_coverage_and_no_fake_evidence(self):
        path = self.root / "empty-text.xlsx"
        workbook(path, [["e", None, None, None, None, None, "图片说它很美", None, None]])
        catalogue, _ = p.inventory(path)
        self.assertEqual(catalogue["evidence"], [])
        self.assertEqual(catalogue["articles"][0]["evidence_ids"], [])

    def test_multiple_batches_merge_without_losing_sources(self):
        source = self.root / "longer.xlsx"
        book = load_workbook(self.source)
        for row in (2, 3, 4):
            book.active.cell(row, 4).value += "背景材料。" * 150
        book.save(source)
        p.prepare(source, self.root / "multi-preview")
        preview = p.read(self.root / "multi-preview/batches/B0001/request.json")
        payload = json.loads(preview["messages"][1]["content"])["articles"]
        prompt = preview["messages"][0]["content"]
        max_chars = max(sum(len(msg["content"]) for msg in p.packet(prompt, [item], "articles")) for item in payload)
        run = self.root / "multi"
        manifest = p.prepare(source, run, max_chars=max_chars)
        self.assertEqual(len(manifest["batch_ids"]), 3)
        children = []
        original = self.proposal()
        for bid in manifest["batch_ids"]:
            request = p.read(run / "batches" / bid / "request.json")
            proposal = copy.deepcopy(original)
            proposal.update(batch_id=bid, request_sha256=request["request_sha256"])
            proposal["coverage"] = [row for row in proposal["coverage"] if row["source_id"] in request["source_ids"]]
            proposal["clusters"] = [row for row in proposal["clusters"] if row["members"][0]["evidence_id"].rsplit(":", 1)[0] in request["source_ids"]]
            for row in proposal["clusters"]:
                row["cluster_id"] = bid + "-C01"
                for index, member in enumerate(row["members"]):
                    member["point_id"] = f"{bid}-P{index}"
                children.append(row["cluster_id"])
            p.write(self.root / "part.json", proposal)
            p.check_batch(run, bid, self.root / "part.json")
        p.collect(run)
        cluster = theme("TC01")
        cluster["candidate_ids"] = children
        p.write(self.root / "multi-plan.json", {"collection_sha256": p.digest(run / "collection.json"), "analysis": ANALYSIS, "clusters": [cluster]})
        p.finalize(run, self.root / "multi-plan.json", self.root / "multi-final")
        self.assertEqual(len(p.read(self.root / "multi-final/coverage.json")), 3)

    def test_no_aesthetic_candidates_remains_valid_empty_result(self):
        proposal = self.proposal()
        proposal["clusters"] = []
        for row in proposal["coverage"]:
            row["status"] = "insufficient_text"
            row["reason"] = "合成边界用例：本轮没有足够依据形成候选"
        p.write(self.root / "empty-proposal.json", proposal)
        p.check_batch(self.run, "B0001", self.root / "empty-proposal.json")
        p.collect(self.run)
        p.write(self.root / "empty-plan.json", {"collection_sha256": p.digest(self.run / "collection.json"), "analysis": ANALYSIS, "clusters": []})
        p.finalize(self.run, self.root / "empty-plan.json", self.root / "empty-final")
        self.assertEqual(p.read(self.root / "empty-final/trend-clusters.json")["clusters"], [])
        self.assertEqual(len(p.read(self.root / "empty-final/coverage.json")), 3)

    def test_estimate_matches_frozen_requests(self):
        estimate = p.read(self.run / "cost-plan.json")
        self.assertEqual(estimate["map_packets"], 1)
        self.assertEqual(estimate["map_input_characters"], sum(len(m["content"]) for m in self.request["messages"]))
        self.assertIsNone(estimate["token_estimate"])
        with self.assertRaises(ValueError):
            p.prepare(self.source, self.root / "negative", max_review_calls=-1)
        self.assertFalse((self.root / "negative").exists())

    def test_point_ids_and_fields_required(self):
        for key in ("point_id", "subject", "feature", "dimension", "condition"):
            proposal = self.proposal()
            proposal["clusters"][0]["members"][0].pop(key)
            with self.subTest(key=key), self.assertRaises(ValueError):
                p.validate_batch(proposal, self.request, self.catalogue)
        proposal = self.proposal()
        proposal["clusters"][1]["members"][0]["point_id"] = "B0001-P1"
        with self.assertRaises(ValueError):
            p.validate_batch(proposal, self.request, self.catalogue)

    def test_one_quote_can_support_distinct_points_but_not_duplicate_mechanism(self):
        proposal = self.proposal()
        member = copy.deepcopy(proposal["clusters"][0]["members"][0])
        member.update(point_id="B0001-P3", subject="边缘", feature="局部细节位置", dimension="detail")
        proposal["clusters"][0]["members"].append(member)
        p.validate_batch(proposal, self.request, self.catalogue)
        duplicate = copy.deepcopy(member)
        duplicate["point_id"] = "B0001-P4"
        proposal["clusters"][0]["members"].append(duplicate)
        with self.assertRaisesRegex(ValueError, "要点重复"):
            p.validate_batch(proposal, self.request, self.catalogue)

    def test_split_without_replaying_batch_and_reject_lost_or_duplicate_point(self):
        proposal = self.proposal()
        proposal["clusters"][0]["members"].extend(proposal["clusters"].pop()["members"])
        p.write(self.root / "proposal.json", proposal)
        p.check_batch(self.run, "B0001", self.root / "proposal.json")
        p.collect(self.run)
        clusters = []
        for index in (1, 2):
            cluster = theme(f"TC{index}")
            cluster.update(candidate_ids=["B0001-C1"], point_ids=[f"B0001-P{index}"])
            clusters.append(cluster)
        plan = {"collection_sha256": p.digest(self.run / "collection.json"), "analysis": ANALYSIS, "clusters": clusters}
        path = self.root / "split-plan.json"
        p.write(path, plan)
        p.finalize(self.run, path, self.root / "split-final")
        result = p.read(self.root / "split-final/trend-clusters.json")["clusters"]
        self.assertEqual([c["source_ids"] for c in result], [["T:1"], ["T:2"]])
        self.assertEqual([len(c["lineage"][0]["members"]) for c in result], [1, 1])
        for selections in (["B0001-P1"], ["B0001-P9"]):
            changed = copy.deepcopy(plan)
            changed["clusters"][1]["point_ids"] = selections
            p.write(path, changed)
            with self.assertRaises(ValueError):
                p.compile_plan(self.run, path)
        plan["clusters"].pop()
        p.write(path, plan)
        with self.assertRaisesRegex(ValueError, "遗漏"):
            p.compile_plan(self.run, path)

    def test_representatives_keep_conditions_diversity_and_full_negation(self):
        evidence = {str(i): {"source_group_id": "same" if i < 3 else f"G{i}",
                              "text": "主体不采用镜面，只在边框保留高光。"} for i in range(5)}
        members = [{"point_id": f"P{i}", "evidence_id": str(i), "subject": "主体" if i < 3 else "边框",
                    "dimension": "surface", "condition": "仅局部" if i == 4 else None, "quote": "镜面"} for i in range(5)]
        result = p.representative_points(members, evidence)
        self.assertEqual(result[0]["point_id"], "P4")
        self.assertEqual(len({evidence[m["evidence_id"]]["source_group_id"] for m in result}), 3)
        self.assertTrue(all("不采用镜面" in row["context"] for row in result))
        self.assertEqual(result, p.representative_points(list(reversed(members)), evidence))

    def test_review_cap_zero_overflow_and_plan_binding(self):
        plan_path = self.accept_and_collect()
        manifest, catalogue, _, _, clusters = p.compile_plan(self.run, plan_path)
        cluster = clusters[0]
        cluster["review_reasons"] = ["审查共同机制"]
        copies = [{**cluster, "cluster_id": f"TC{i}"} for i in range(5)]
        queue = p.review_packets(manifest, catalogue, copies)
        self.assertEqual(len(queue["packets"]), 3)
        self.assertEqual(len(queue["deferred"]), 2)
        self.assertFalse(queue["review_performed"])
        zero = p.review_packets({**manifest, "max_review_calls": 0}, catalogue, copies)
        self.assertEqual(len(zero["packets"]), 0)
        self.assertEqual(len(zero["deferred"]), 5)
        tiny = p.review_packets({**manifest, "max_chars": 20}, catalogue, copies)
        self.assertEqual(len(tiny["deferred"]), 5)
        plan = p.read(plan_path)
        plan["clusters"][0]["review_reasons"] = ["审查共同机制"]
        p.write(plan_path, plan)
        out = self.root / "review.json"
        p.review(self.run, plan_path, out)
        self.assertEqual(p.read(out)["plan_sha256"], p.digest(plan_path))
        with self.assertRaises(ValueError):
            p.review(self.run, plan_path, out)


if __name__ == "__main__":
    unittest.main()
