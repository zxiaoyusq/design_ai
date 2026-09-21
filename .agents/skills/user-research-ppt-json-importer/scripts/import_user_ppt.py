#!/usr/bin/env python3
"""依据人工审核后的 manifest，将单个用户 PPT 的问答和图片补录到研究 JSON。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ALLOWED_EMOTIONS = {"ENJOY", "DISLIKE", "REFERENCE"}
RESERVED_LINK_FIELDS = {
    "code",
    "image_ids",
    "status",
    "local_paths",
    "emotion_tag",
    "source_type",
    "source_slide",
    "evidence_scope",
}


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def slugify(value: str) -> str:
    value = value.lower().replace("–", "-")
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value or "image"


def shape_at(slide, shape_path: str):
    indexes = [int(part) for part in shape_path.split("/") if part]
    if not indexes or not shape_path.startswith("/"):
        raise ValueError(f"无效 shape_path：{shape_path}")
    shape = slide.shapes[indexes[0]]
    for index in indexes[1:]:
        if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
            raise ValueError(f"shape_path 穿过非组合形状：{shape_path}")
        shape = shape.shapes[index]
    return shape


def validated_image_subdir(value: str) -> PurePosixPath:
    path = PurePosixPath(value)
    if path.is_absolute() or ".." in path.parts or not path.parts or path.parts[0] != "ppt_images":
        raise ValueError("image_subdir 必须是位于 ppt_images/ 下的相对路径")
    if len(path.parts) < 2:
        raise ValueError("image_subdir 必须包含独立的用户子目录")
    return path


def validate_manifest(manifest: dict[str, Any], slide_count: int) -> None:
    required = {"source_file", "source_language", "translated_to", "image_subdir", "records"}
    missing = required - set(manifest)
    if missing:
        raise ValueError(f"manifest 缺少字段：{sorted(missing)}")
    validated_image_subdir(manifest["image_subdir"])
    if not isinstance(manifest["records"], list) or not manifest["records"]:
        raise ValueError("records 必须是非空数组")

    record_ids: set[str] = set()
    picture_locations: set[tuple[int, str]] = set()
    for record in manifest["records"]:
        for field in (
            "id",
            "source_slide",
            "scenario_type",
            "question_type",
            "question",
            "answer",
        ):
            if not record.get(field):
                raise ValueError(f"record 缺少非空字段 {field}：{record.get('id')}")
        if record["id"] in record_ids:
            raise ValueError(f"record id 重复：{record['id']}")
        record_ids.add(record["id"])
        slide = int(record["source_slide"])
        if not 1 <= slide <= slide_count:
            raise ValueError(f"source_slide 超出范围：{slide}")

        for image in record.get("images", []):
            for field in ("shape_path", "code", "emotion_tag"):
                if not image.get(field):
                    raise ValueError(f"图片条目缺少 {field}：{record['id']}")
            if image["emotion_tag"] not in ALLOWED_EMOTIONS:
                raise ValueError(f"无效 emotion_tag：{image['emotion_tag']}")
            location = (slide, image["shape_path"])
            if location in picture_locations:
                raise ValueError(f"同一 PPT 图片被重复映射：第 {slide} 页 {image['shape_path']}")
            picture_locations.add(location)
            metadata = image.get("link_metadata", {})
            if RESERVED_LINK_FIELDS & set(metadata):
                raise ValueError(f"link_metadata 覆盖保留字段：{record['id']} {image['code']}")

        for image in record.get("missing_images", []):
            for field in ("code", "emotion_tag"):
                if not image.get(field):
                    raise ValueError(f"缺图条目缺少 {field}：{record['id']}")
            if image["emotion_tag"] not in ALLOWED_EMOTIONS:
                raise ValueError(f"无效 emotion_tag：{image['emotion_tag']}")


def qa_record(record: dict[str, Any], manifest: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": record["id"],
        "scenario_type": record["scenario_type"],
        "question_type": record["question_type"],
        "question": record["question"],
        "ai_analysis": record["answer"],
        "bid": record["id"],
        "answer_type": record.get("answer_type", "用户原始填写（PPT翻译补录）"),
        "source_type": "ppt_supplement",
        "source_file": manifest["source_file"],
        "source_slide": int(record["source_slide"]),
        "source_language": manifest["source_language"],
        "translated_to": manifest["translated_to"],
    }


def preference_evidence(image: dict[str, Any], record_id: str, preference_code: str) -> dict[str, Any]:
    scope = image["evidence_scope"]
    if preference_code != image["ref_pic_code"]:
        scope = "group_contains_code"
    return {
        "image_id": image["image_id"],
        "local_path": image.get("local_path"),
        "status": image["status"],
        "source_type": "ppt_supplement",
        "source_file": image["source_file"],
        "source_slide": image["source_slide"],
        "source_question_id": record_id,
        "evidence_scope": scope,
        "source_ref_pic_code": image["ref_pic_code"],
    }


def recompute_counts(data: dict[str, Any]) -> dict[str, int]:
    research = [item for user in data["users"] for item in user.get("aesthetic_research", [])]
    demand = [item for user in data["users"] for item in user.get("demand_research", [])]
    preferences = [item for user in data["users"] for item in user.get("image_preferences", [])]
    return {
        "users": len(data["users"]),
        "aesthetic_research": len(research),
        "demand_research": len(demand),
        "image_preferences": len(preferences),
        "enjoy": sum(item.get("emotion_tag") == "ENJOY" for item in preferences),
        "dislike": sum(item.get("emotion_tag") == "DISLIKE" for item in preferences),
        "ppt_supplement_aesthetic_research": sum(
            item.get("source_type") == "ppt_supplement" for item in research
        ),
        "ppt_supplement_image_preferences": sum(
            item.get("source_type") == "ppt_supplement" for item in preferences
        ),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ppt", required=True, type=Path)
    parser.add_argument("--json", required=True, type=Path)
    parser.add_argument("--user-name", required=True)
    parser.add_argument("--manifest", required=True, type=Path)
    args = parser.parse_args()

    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    data = json.loads(args.json.read_text(encoding="utf-8"))
    presentation = Presentation(args.ppt)
    validate_manifest(manifest, len(presentation.slides))

    matching_users = [user for user in data.get("users", []) if user.get("name") == args.user_name]
    if len(matching_users) != 1:
        raise ValueError(f"目标用户名必须精确匹配一条记录，实际匹配 {len(matching_users)} 条")
    user = matching_users[0]
    source_file = manifest["source_file"]
    image_subdir = validated_image_subdir(manifest["image_subdir"])
    output_dir = args.json.parent.joinpath(*image_subdir.parts)
    staging_dir = output_dir.parent / f".{output_dir.name}-staging"
    backup_dir = output_dir.parent / f".{output_dir.name}-backup"

    # 清理仅限同一来源，确保重新执行不会影响其他 PPT 或原始研究数据。
    base_research = [
        item
        for item in user.get("aesthetic_research", [])
        if not (
            item.get("source_type") == "ppt_supplement"
            and item.get("source_file") == source_file
        )
    ]
    base_preferences = []
    for item in user.get("image_preferences", []):
        if item.get("source_type") == "ppt_supplement" and item.get("source_file") == source_file:
            continue
        preference = dict(item)
        if preference.get("ppt_evidence"):
            retained = [
                evidence
                for evidence in preference["ppt_evidence"]
                if not (
                    evidence.get("source_type") == "ppt_supplement"
                    and evidence.get("source_file") == source_file
                )
            ]
            if retained:
                preference["ppt_evidence"] = retained
            else:
                preference.pop("ppt_evidence", None)
        base_preferences.append(preference)
    base_ppt_images = [
        item for item in user.get("ppt_images", []) if item.get("source_file") != source_file
    ]

    if staging_dir.exists():
        shutil.rmtree(staging_dir)
    if backup_dir.exists():
        raise RuntimeError(f"检测到未处理的备份目录，请先检查：{backup_dir}")
    staging_dir.mkdir(parents=True, exist_ok=True)

    ppt_sha256 = sha256_bytes(args.ppt.read_bytes())
    manifest_sha256 = sha256_bytes(args.manifest.read_bytes())
    image_records: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []

    try:
        for record_spec in manifest["records"]:
            record = qa_record(record_spec, manifest)
            linked_images: list[dict[str, Any]] = []
            slide_number = int(record_spec["source_slide"])
            slide = presentation.slides[slide_number - 1]

            for image_spec in record_spec.get("images", []):
                shape = shape_at(slide, image_spec["shape_path"])
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    raise ValueError(
                        f"形状不是图片：第 {slide_number} 页 {image_spec['shape_path']}"
                    )
                blob = shape.image.blob
                digest = sha256_bytes(blob)
                extension = (shape.image.ext or "bin").lower().replace("jpg", "jpeg")
                image_id = "ppt_image_" + hashlib.sha256(
                    (
                        f"{ppt_sha256}:{slide_number}:{image_spec['shape_path']}:"
                        f"{image_spec['code']}:{digest}"
                    ).encode("utf-8")
                ).hexdigest()[:24]
                filename = (
                    f"s{slide_number:02d}_{slugify(image_spec['code'])}_"
                    f"{image_id[-8:]}.{extension}"
                )
                staged_path = staging_dir / filename
                staged_path.write_bytes(blob)
                with Image.open(staged_path) as image:
                    image.verify()
                with Image.open(staged_path) as image:
                    width, height = image.size
                    image_format = image.format

                image_record = {
                    "image_id": image_id,
                    "ref_pic_code": image_spec["code"],
                    "local_path": (image_subdir / filename).as_posix(),
                    "emotion_tag": image_spec["emotion_tag"],
                    "status": "extracted",
                    "source_type": "ppt_supplement",
                    "source_file": source_file,
                    "source_slide": slide_number,
                    "source_shape_path": image_spec["shape_path"],
                    "source_question_id": record_spec["id"],
                    "evidence_scope": image_spec.get("evidence_scope", "exact"),
                    "description_zh": image_spec.get("description_zh", ""),
                    "sha256": digest,
                    "format": image_format,
                    "width": width,
                    "height": height,
                    "bytes": len(blob),
                }
                if image_spec.get("contained_codes"):
                    image_record["contained_codes"] = image_spec["contained_codes"]
                image_record["_preference_codes"] = image_spec.get(
                    "preference_codes",
                    [image_spec["code"]]
                    if image_spec["emotion_tag"] in {"ENJOY", "DISLIKE"}
                    else [],
                )
                image_record["_link_metadata"] = image_spec.get("link_metadata", {})
                image_records.append(image_record)
                linked_images.append(image_record)

            for image_spec in record_spec.get("missing_images", []):
                image_id = "ppt_missing_" + hashlib.sha256(
                    f"{ppt_sha256}:{slide_number}:{record_spec['id']}:{image_spec['code']}".encode(
                        "utf-8"
                    )
                ).hexdigest()[:24]
                image_record = {
                    "image_id": image_id,
                    "ref_pic_code": image_spec["code"],
                    "local_path": None,
                    "emotion_tag": image_spec["emotion_tag"],
                    "status": "missing_source_image",
                    "source_type": "ppt_supplement",
                    "source_file": source_file,
                    "source_slide": slide_number,
                    "source_shape_path": None,
                    "source_question_id": record_spec["id"],
                    "evidence_scope": image_spec.get("evidence_scope", "declared_code_only"),
                    "description_zh": image_spec.get("description_zh", ""),
                    "_preference_codes": image_spec.get(
                        "preference_codes",
                        [image_spec["code"]]
                        if image_spec["emotion_tag"] in {"ENJOY", "DISLIKE"}
                        else [],
                    ),
                    "_link_metadata": image_spec.get("link_metadata", {}),
                }
                image_records.append(image_record)
                linked_images.append(image_record)

            if linked_images:
                record["ref_pic"] = ",".join(item["ref_pic_code"] for item in linked_images)
                record["ref_pic_links"] = []
                for item in linked_images:
                    link = {
                        "code": item["ref_pic_code"],
                        "image_ids": [item["image_id"]] if item["status"] == "extracted" else [],
                        "status": "matched" if item["status"] == "extracted" else item["status"],
                        "local_paths": [item["local_path"]] if item.get("local_path") else [],
                        "emotion_tag": item["emotion_tag"],
                        "source_type": "ppt_supplement",
                        "source_slide": item["source_slide"],
                        "evidence_scope": item["evidence_scope"],
                    }
                    if item.get("contained_codes"):
                        link["contained_codes"] = item["contained_codes"]
                    link.update(item["_link_metadata"])
                    record["ref_pic_links"].append(link)
            records.append(record)

        extracted_count = sum(item["status"] == "extracted" for item in image_records)
        if len(list(staging_dir.iterdir())) != extracted_count:
            raise RuntimeError("暂存目录图片数量与提取记录不一致")

        preference_index = {
            (item.get("ref_pic_code"), item.get("emotion_tag")): item
            for item in base_preferences
        }
        added_preference_codes: list[str] = []
        for image in image_records:
            for preference_code in image["_preference_codes"]:
                emotion = image["emotion_tag"]
                if emotion not in {"ENJOY", "DISLIKE"}:
                    raise ValueError("REFERENCE 图片不能生成 image_preferences")
                key = (preference_code, emotion)
                evidence = preference_evidence(image, image["source_question_id"], preference_code)
                if key in preference_index:
                    existing_evidence = preference_index[key].setdefault("ppt_evidence", [])
                    if evidence not in existing_evidence:
                        existing_evidence.append(evidence)
                    continue

                preference_id = "ppt_pref_" + hashlib.sha256(
                    f"{ppt_sha256}:{preference_code}:{emotion}".encode("utf-8")
                ).hexdigest()[:24]
                preference = {
                    "image_id": preference_id,
                    "source_id": image["source_question_id"],
                    "source_bid": image["source_question_id"],
                    "emotion_tag": emotion,
                    "ref_pic_code": preference_code,
                    "local_path": image.get("local_path"),
                    "status": image["status"],
                    "source_type": "ppt_supplement",
                    "source_file": source_file,
                    "source_slide": image["source_slide"],
                    "ppt_image_id": image["image_id"] if image["status"] == "extracted" else None,
                    "evidence_scope": (
                        image["evidence_scope"]
                        if preference_code == image["ref_pic_code"]
                        else "group_contains_code"
                    ),
                }
                if preference_code != image["ref_pic_code"]:
                    preference["source_ref_pic_code"] = image["ref_pic_code"]
                base_preferences.append(preference)
                preference_index[key] = preference
                added_preference_codes.append(preference_code)

        clean_image_records = []
        for item in image_records:
            clean = dict(item)
            clean.pop("_preference_codes", None)
            clean.pop("_link_metadata", None)
            clean_image_records.append(clean)

        user["aesthetic_research"] = base_research + records
        user["image_preferences"] = base_preferences
        user["ppt_images"] = base_ppt_images + clean_image_records

        now = datetime.now(timezone.utc).isoformat()
        old_supplement = user.get("ppt_supplement", {})
        old_added_at = (
            old_supplement.get("added_at")
            if old_supplement.get("source_file") == source_file
            else None
        )
        user["ppt_supplement"] = {
            "source_type": "ppt_supplement",
            "source_file": source_file,
            "source_sha256": ppt_sha256,
            "manifest_sha256": manifest_sha256,
            "source_language": manifest["source_language"],
            "translated_to": manifest["translated_to"],
            "added_at": old_added_at or now,
            "updated_at": now,
            "inclusion_rule": "仅补录 PPT 中同时存在问题和用户回答的内容；图片选择可以作为回答，只有问题的主持人提示不录入。",
            "aesthetic_research_ids": [item["id"] for item in records],
            "image_directory": image_subdir.as_posix(),
            "extracted_image_count": extracted_count,
            "missing_source_image_count": sum(
                item["status"] == "missing_source_image" for item in clean_image_records
            ),
            "image_preference_codes_added": added_preference_codes,
        }
        data["counts"] = recompute_counts(data)

        if any(not item.get("question") or not item.get("ai_analysis") for item in records):
            raise RuntimeError("存在只有问题或只有回答的补录记录")
        for item in clean_image_records:
            if item["status"] == "extracted":
                file_path = staging_dir / Path(item["local_path"]).name
                if not file_path.exists():
                    raise RuntimeError(f"提取文件不存在：{file_path}")

        temporary_json = args.json.with_suffix(args.json.suffix + ".tmp")
        temporary_json.write_text(
            json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        json.loads(temporary_json.read_text(encoding="utf-8"))

        # 目录和 JSON 一起提交；任何一步失败都恢复旧图片目录。
        if output_dir.exists():
            os.replace(output_dir, backup_dir)
        try:
            os.replace(staging_dir, output_dir)
            os.replace(temporary_json, args.json)
        except Exception:
            if output_dir.exists():
                shutil.rmtree(output_dir)
            if backup_dir.exists():
                os.replace(backup_dir, output_dir)
            raise
        if backup_dir.exists():
            shutil.rmtree(backup_dir)

        print(
            json.dumps(
                {
                    "user": args.user_name,
                    "records_added": len(records),
                    "image_evidence": len(clean_image_records),
                    "images_extracted": extracted_count,
                    "missing_source_images": sum(
                        item["status"] == "missing_source_image"
                        for item in clean_image_records
                    ),
                    "image_emotions": Counter(
                        item["emotion_tag"] for item in clean_image_records
                    ),
                    "preferences_added": len(added_preference_codes),
                    "counts": data["counts"],
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    except Exception:
        if staging_dir.exists():
            shutil.rmtree(staging_dir)
        raise


if __name__ == "__main__":
    main()
