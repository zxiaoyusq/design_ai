#!/usr/bin/env python3
"""输出 PPT 的逐页文本、备注和内嵌图片路径，供语义审核和 manifest 编写。"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def collect_shape(shape, path: str, pictures: list[dict], texts: list[dict]) -> None:
    text = getattr(shape, "text", "").strip()
    if text:
        texts.append(
            {
                "shape_path": path,
                "text": text,
                "left": int(shape.left),
                "top": int(shape.top),
                "width": int(shape.width),
                "height": int(shape.height),
            }
        )
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        blob = shape.image.blob
        pictures.append(
            {
                "shape_path": path,
                "name": shape.name,
                "left": int(shape.left),
                "top": int(shape.top),
                "width": int(shape.width),
                "height": int(shape.height),
                "extension": shape.image.ext,
                "sha256": hashlib.sha256(blob).hexdigest(),
                "bytes": len(blob),
            }
        )
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for index, child in enumerate(shape.shapes):
            collect_shape(child, f"{path}/{index}", pictures, texts)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ppt", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    presentation = Presentation(args.ppt)
    slides = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        pictures: list[dict] = []
        texts: list[dict] = []
        for index, shape in enumerate(slide.shapes):
            collect_shape(shape, f"/{index}", pictures, texts)

        notes = []
        try:
            for shape in slide.notes_slide.shapes:
                text = getattr(shape, "text", "").strip()
                if text:
                    notes.append(text)
        except (AttributeError, ValueError):
            pass

        slides.append(
            {
                "slide": slide_number,
                "texts": texts,
                "notes": notes,
                "pictures": pictures,
            }
        )

    result = {
        "ppt": str(args.ppt),
        "ppt_sha256": hashlib.sha256(args.ppt.read_bytes()).hexdigest(),
        "slide_count": len(slides),
        "slides": slides,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"slides": len(slides), "pictures": sum(len(item["pictures"]) for item in slides)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
