"""将原 HTML 截图铺满 16:9 PPTX，可为第四页添加不改画面的原图放大入口。"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


def add_click_region(slide, bounds: list[int], target, name: str):
    """实心但完全透明的形状保留整块命中区，避免无填充形状只能点击边框。"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *bounds)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    alpha = OxmlElement('a:alpha')
    alpha.set('val', '0')
    shape._element.spPr.solidFill.srgbClr.append(alpha)
    shape.line.fill.background()
    shape.click_action.target_slide = target
    return shape


def add_text(slide, text: str, x: float, y: float, w: float, h: float, size: int, color: str):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = 'Noto Sans CJK SC'
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = RGBColor.from_string(color)
    return shape


def add_zoom_slide(deck, metadata: dict, capture_dir: Path, origin):
    """追加隐藏的原图页；顺序放映跳过，图片与返回按钮均能回到第四页。"""
    path = capture_dir / metadata['filename']
    assert hashlib.sha256(path.read_bytes()).hexdigest() == metadata['sha256']
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.name = f'第 4 页放大 · {metadata["title"]}'
    slide._element.set('show', '0')
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('071C2C')
    title = '真实素材 · 原图放大' if metadata['key'] == 'dnaimage' else '专业关系 · DNA 分析放大'
    add_text(slide, title, .36, .12, 9, .38, 18, 'F4F7F9')
    add_text(slide, '点击图片或右上角按钮返回第 4 页', .36, 7.17, 11, .22, 11, '9AB2BF')
    # 保留原图字节与完整比例；不把缩略图再放大，不裁掉结果文字。
    available_w, available_h = Inches(12.61), Inches(6.48)
    scale = min(available_w / metadata['image_width'], available_h / metadata['image_height'])
    width, height = round(metadata['image_width'] * scale), round(metadata['image_height'] * scale)
    left = (deck.slide_width - width) // 2
    top = Inches(.60) + (available_h - height) // 2
    image = slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    image.name = f'原图 · {metadata["key"]}'
    image.click_action.target_slide = origin
    back = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.92), Inches(.12), Inches(2.06), Inches(.38))
    back.name = '返回第 4 页'
    back.fill.solid()
    back.fill.fore_color.rgb = RGBColor.from_string('12364B')
    back.line.color.rgb = RGBColor.from_string('4ED9CF')
    back.text_frame.text = '返回第 4 页  ×'
    back.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = back.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.name = 'Noto Sans CJK SC'
    paragraph.font.size = Pt(12)
    paragraph.font.color.rgb = RGBColor.from_string('4ED9CF')
    back.click_action.target_slide = origin
    # 放大页不靠空白点击顺序前进，防止意外进入下一张隐藏图。
    transition = OxmlElement('p:transition')
    transition.set('advClick', '0')
    slide._element.append(transition)
    slide.notes_slide.notes_text_frame.text = f'第 4 页图片放大：{metadata["title"]}。仅原图展示，不新增分析结论。点击图片或右上角返回第 4 页；本页隐藏，不计入主讲或备答。'
    return slide


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("capture_dir", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument('--zoom-slide4', action='store_true', help='为第四页两张图添加透明点击区、隐藏放大页与返回链接')
    args = parser.parse_args()
    capture = json.loads((args.capture_dir / "capture.json").read_text())
    source = Path(capture["source"])
    assert hashlib.sha256(source.read_bytes()).hexdigest() == capture["source_sha256"]

    deck = Presentation()
    # 使用精确 EMU 尺寸，保证每页图片无裁切、无留边地铺满 16:9 画布。
    deck.slide_width, deck.slide_height = 12_192_000, 6_858_000
    deck.core_properties.title = "棱镜：让每一次设计选择，更贴近用户期待"
    deck.core_properties.subject = "原 HTML 逐页高清截图版"
    deck.core_properties.author = "PRISM 项目组"
    deck.core_properties.keywords = "PRISM, 棱镜, 产品路演, 原页面截图"
    deck.core_properties.comments = "按原 HTML 顺序生成，每页一张完整截图；可选第四页原图放大，不依赖外部资源。"
    for item in capture["slides"]:
        image = args.capture_dir / item["filename"]
        assert hashlib.sha256(image.read_bytes()).hexdigest() == item["sha256"]
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.name = f'{item["number"]:02} {item["title"]}'
        picture = slide.shapes.add_picture(str(image), 0, 0, width=deck.slide_width, height=deck.slide_height)
        picture.name = item["title"]
        picture._element.nvPicPr.cNvPr.set("descr", item["title"])
    zoom_targets = []
    if args.zoom_slide4:
        assert len(capture['slides']) >= 4
        hotspots = capture['slides'][3].get('zoom_hotspots', [])
        assert [entry['key'] for entry in hotspots] == ['dnaimage', 'dnarelations'], '请先用当前截图脚本记录第四页两张图片的真实位置'
        origin = deck.slides[3]
        for metadata in hotspots:
            target = add_zoom_slide(deck, metadata, args.capture_dir, origin)
            x, y, w, h = metadata['bounds']
            assert 0 <= x < x+w <= 1 and 0 <= y < y+h <= 1
            bounds = [round(x * deck.slide_width), round(y * deck.slide_height), round(w * deck.slide_width), round(h * deck.slide_height)]
            add_click_region(origin, bounds, target, f'点击放大 · {metadata["key"]}')
            zoom_targets.append({'key': metadata['key'], 'slide': len(deck.slides), 'bounds_emu': bounds, 'source_sha256': metadata['sha256']})
        origin.notes_slide.notes_text_frame.text = '放映模式下点击“真实素材”或“专业关系”图片打开原图放大页，点击放大图片或右上角按钮返回第 4 页。两张放大页隐藏，正常翻页保持原 13 页顺序。'
    deck.save(args.output)

    # 复读实际交付，核对页数、全幅坐标与图片字节；避免遗漏备用页或图片被重新压缩。
    actual = Presentation(args.output)
    assert len(actual.slides) == len(capture["slides"]) + len(zoom_targets)
    for slide, item in zip(actual.slides, capture["slides"]):
        assert len(slide.shapes) == (3 if args.zoom_slide4 and item['number'] == 4 else 1)
        assert slide._element.get('show') != '0'
        image = slide.shapes[0]
        assert image.shape_type == MSO_SHAPE_TYPE.PICTURE
        assert (image.left, image.top, image.width, image.height) == (0, 0, actual.slide_width, actual.slide_height)
        assert (image.crop_left, image.crop_top, image.crop_right, image.crop_bottom) == (0, 0, 0, 0)
        assert image.image.size == (3200, 1800)
        assert hashlib.sha256(image.image.blob).hexdigest() == item["sha256"]
    for index, metadata in enumerate(zoom_targets):
        target = actual.slides[metadata['slide'] - 1]
        hotspot = actual.slides[3].shapes[index + 1]
        assert target._element.get('show') == '0'
        assert hotspot.click_action.target_slide == target
        assert hotspot._element.spPr.solidFill.srgbClr[0].get('val') == '0'
        assert [hotspot.left, hotspot.top, hotspot.width, hotspot.height] == metadata['bounds_emu']
        pictures = [shape for shape in target.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pictures) == 1
        assert hashlib.sha256(pictures[0].image.blob).hexdigest() == metadata['source_sha256']
        assert all(crop == 0 for crop in (pictures[0].crop_left, pictures[0].crop_top, pictures[0].crop_right, pictures[0].crop_bottom))
        assert pictures[0].click_action.target_slide == actual.slides[3]
        assert next(shape for shape in target.shapes if shape.name == '返回第 4 页').click_action.target_slide == actual.slides[3]
    with zipfile.ZipFile(args.output) as archive:
        assert archive.testzip() is None
    receipt = {
        "pptx": args.output.name,
        "pptx_bytes": args.output.stat().st_size,
        "pptx_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest(),
        "source_sha256": capture["source_sha256"],
        "slide_count": len(actual.slides),
        "visible_slide_count": len(capture['slides']),
        "hidden_zoom_slide_count": len(zoom_targets),
        "zoom_links": zoom_targets,
        "main_slide_count": sum(not item["backup"] for item in capture["slides"]),
        "backup_slide_count": sum(item["backup"] for item in capture["slides"]),
        "aspect_ratio": "16:9",
        "screenshot_pixels": [3200, 1800],
        "checks": "可见页顺序、全幅无裁切、截图原字节、原图字节、放大页隐藏状态、点击位置、双向跳转和 ZIP 完整性通过",
    }
    (args.capture_dir / "export.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(receipt, ensure_ascii=False))


if __name__ == "__main__":
    main()
