"""生成领导汇报版 HTML 与同源讲稿；保留原图字节，主讲和备用时长分开。"""

from __future__ import annotations

import base64
import argparse
import hashlib
import html
import json
from pathlib import Path
import re
import zipfile


def stamp(seconds: int) -> str:
    return f"{seconds // 60:02}:{seconds % 60:02}"


def label(index: int, main_count: int) -> str:
    return f"{index + 1:02}" if index < main_count else f"B{index - main_count + 1}"


def timings(talk: dict) -> list[str]:
    total = 0
    result = []
    for index, slide in enumerate(talk["slides"]):
        if index < talk["mainCount"]:
            result.append(f"{stamp(total)}–{stamp(total + slide['seconds'])} · 建议 {slide['seconds']} 秒")
            total += slide["seconds"]
        else:
            result.append(f"按问题选用 · 建议 {slide['seconds']} 秒 · 不计入主讲")
    return result


def make_markdown(talk: dict, spoken_chars: int) -> str:
    main_count = talk["mainCount"]
    lines = [
        f"# {talk['title']}", "",
        "PRISM 棱镜｜领导汇报演讲稿 · 2026 年 10 月 26 日", "",
        f"{main_count} 页主讲，约 12 分钟；4 页备用，按问题选用；预留约 3 分钟问答，总体控制在 15 分钟内。", "",
        f"主讲正文与转场共约 {spoken_chars:,} 字符。按每分钟约 220–240 字符排练，主讲包含约 15 秒猜图互动，并留看图与停顿时间；实际时长以试讲校准。只读口播正文和转场，不朗读标题、页面答案、讲者提示及来源。备用页不全部连讲，每题约 40 秒。", "",
        "## 讲者使用说明", "",
        "主线：双定位 → 多部门用研与需求衰减 → 前后流程与愿景 → 两类挑战与两个故事 → 一期成果与二期规划 → 猜图 → 揭晓与 Camon → 情绪板生成过程 → 用户价值。第 2 页曲线为定性风险示意，不报衰减数值；第 3 页上下对照完整愿景；第 4 页可放大 DNA；第 5 页强调已完成范围；第 6 页猜图、第 7 页揭晓 A 为 AI、B 为设计师；第 8 页按从左到右的五栏讲解。猜图期间不要投屏讲稿或目录。", "",
        f"第 {main_count} 页结束主讲，继续向后翻到第 {main_count + 1}–{len(talk['slides'])} 页（B1–B4）即是备用回答。遇到具体问题，可在目录中直接选相应备用页，无需倒回主内容。", "",
        "方向键、空格翻页，N 打开当前讲稿，F 全屏。讲稿会出现在当前投屏画面，请用这份文档单独查看。计时手动开始，剩余 3 分钟提示问答。", "",
        "## 30 秒核心表达（备用开场）", "", talk["elevator"], "",
        "## 主讲 · 12 分钟", "",
    ]
    for index, (slide, timing) in enumerate(zip(talk["slides"], timings(talk))):
        if index == main_count:
            lines.extend(["## 备用页 · 主内容之后，按问题向后展示", "", "以下不是主讲续篇。选择最相关的问题回答，控制总问答时间；所有回答页均自足，不要求回翻前文。", ""])
        lines.extend([
            f"### {label(index, main_count)}｜{slide['title']}", "", f"建议节奏：{timing}", "",
            f"页面答案：{slide['answer']}", "", f"讲者提示：{slide['cue']}", "",
            *[part for paragraph in slide["paragraphs"] for part in (paragraph, "")],
            *([f"转场：{slide['bridge']}", ""] if slide["bridge"] else []),
        ])
    lines.extend(["## 来源与表达边界（供准备，不逐条宣读）", ""])
    for item in talk["sources"]:
        lines.extend([f"### {item['label']}", "", item["detail"], ""])
    lines.extend([
        "## 排练检查点", "",
        "- 00:00 双定位；00:45 多部门用研与曲线；01:55 前后流程；03:25 两类挑战与两个故事；05:30 一期成果与二期规划。",
        "- 06:50 猜图（含 15 秒互动）；07:30 揭晓与 Camon；08:20 从左到右讲生成过程；10:50 收束；12:00 结束主讲。",
        "- 开场、第 2 页与第 9 页的问题后停顿，再给答案；第 4 页讲故事时放慢；第 6 页留出安静的看图时间。其他页面不逐项念表，避免每句都用问句。",
        "- 如果超时，第 8 页可省略“为什么要分开做”一段；第 5 页二期只概括交付物和人的确认。保留两类挑战、两个故事、MVP 完成状态、猜图揭晓、Camon 与结尾。",
        "- 规模与 3–5 天为项目方提供的汇报口径；周期只指资料已积累后的按需选向、聚类与筛调，不含历史积累、后续用户调研或整机研发。不另推成本金额、用户认可率或畅销概率。",
        "",
    ])
    return "\n".join(lines)


def make_speech_html(talk: dict, spoken_chars: int) -> str:
    sections = []
    toc = []
    for index, (slide, timing) in enumerate(zip(talk["slides"], timings(talk))):
        key = label(index, talk["mainCount"])
        if index == talk["mainCount"]:
            sections.append('<section class="divider"><h2>以下为备用页</h2><p>提问时按需向后展示，不计入 12 分钟主讲。每题约 40 秒，不全部连讲。</p></section>')
            toc.append('<strong class="toc-group">备用 B1–B4 · 按需展示</strong>')
        toc.append(f'<a href="#p{index + 1}">{key} · {html.escape(slide["title"])}</a>')
        body = "".join(f"<p>{html.escape(p)}</p>" for p in slide["paragraphs"])
        bridge = f'<p class="bridge">转场：{html.escape(slide["bridge"])}</p>' if slide["bridge"] else ""
        sections.append(
            f'<section id="p{index + 1}"><div class="timing">{key} · {timing}</div>'
            f'<h2>{html.escape(slide["title"])}</h2><p class="answer">{html.escape(slide["answer"])}</p>'
            f'<aside>{html.escape(slide["cue"])}</aside>{body}{bridge}</section>'
        )
    sources = "".join(f'<h3>{html.escape(item["label"])}</h3><p>{html.escape(item["detail"])}</p>' for item in talk["sources"])
    return f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PRISM 棱镜 · 领导汇报讲稿</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#f4f5f7;color:#17202b;font-family:"Noto Sans CJK SC","PingFang SC","Microsoft YaHei",sans-serif;line-height:1.95}}main{{max-width:860px;padding:55px 35px 90px;margin:auto}}h1{{font-size:40px;line-height:1.4;letter-spacing:-.04em}}h2{{font-size:28px;line-height:1.5;margin:12px 0 18px}}h3{{font-size:19px;margin:24px 0 12px}}p{{font-size:18px;margin:18px 0}}.label,.timing{{font-size:13px;color:#566173}}.answer{{color:#3154ad;font-size:18px}}aside{{padding:14px 18px;background:#e5ebf8;color:#455a8b;font-size:14px}}section{{border-top:1px solid #ccd2dc;margin-top:40px;padding-top:28px;scroll-margin-top:25px}}.bridge{{border-left:3px solid #7d99ef;padding:7px 17px;color:#3454a4;font-size:16px}}.divider{{padding:25px;background:#e2e9f8}}.actions{{display:flex;flex-wrap:wrap;gap:10px;margin:25px 0}}.actions a,.actions button{{font:inherit;font-size:14px;padding:9px 14px;border:1px solid #aebbd3;background:transparent;text-decoration:none;color:#3154ad;border-radius:5px;cursor:pointer}}nav{{display:grid;grid-template-columns:1fr 1fr;gap:10px 20px;margin-top:25px}}nav a{{font-size:13px;color:#3557b1;text-decoration:none;border-bottom:1px solid #ccd2dc;padding:8px 0}}.toc-group{{grid-column:1/-1;font-size:16px;margin-top:17px}}.sources p{{font-size:14px;color:#566173;overflow-wrap:anywhere}}a:focus-visible,button:focus-visible{{outline:3px solid #315cff;outline-offset:4px}}@media(max-width:650px){{main{{padding:30px 22px 65px}}h1{{font-size:31px}}h2{{font-size:25px}}p{{font-size:17px}}nav{{grid-template-columns:1fr}}}}@media print{{@page{{size:A4;margin:18mm}}body{{background:white}}main{{max-width:none;padding:0}}nav,.actions{{display:none}}h1{{font-size:23pt}}h2{{font-size:16pt;break-after:avoid}}h3{{font-size:13pt;break-after:avoid}}p{{font-size:11pt;line-height:1.8}}.timing{{break-after:avoid}}aside{{font-size:9pt}}.answer,.bridge{{font-size:10pt}}}}
</style></head><body><main><div class="label">PRISM · LEADERSHIP BRIEFING · 2026.10.26</div><h1>{html.escape(talk["title"])}</h1><p>{talk["mainCount"]} 页主讲 · 12 分钟<br>4 页独立备答 · 预留 3 分钟按需回答</p><p class="label">主讲约 {spoken_chars:,} 字符，按每分钟约 220–240 字符排练，含约 15 秒互动及看图停顿，实际耗时请试讲校准。只读正文和转场，不读标题、页面答案、提示与来源。</p><div class="actions"><a href="index.html">打开演示页</a><a href="PRISM-产品路演演讲稿.docx" download>下载 Word 稿</a><a href="Untitled.md" download>下载指定文字稿</a><button type="button" onclick="window.print()">打印 / 保存 PDF</button></div><aside>第 {talk["mainCount"]} 页结束主讲；第 {talk["mainCount"] + 1}–{len(talk["slides"])} 页是独立备用回答。按问题选择，不需要回翻前文，也不要将 4 页全部连讲。讲稿弹层会显示在投屏画面，正式汇报请单独查看本文档。</aside><nav aria-label="讲稿目录">{''.join(toc)}</nav><section><h2>30 秒核心表达</h2><p>{html.escape(talk["elevator"])}</p></section>{''.join(sections)}<section class="sources"><h2>来源与表达边界</h2>{sources}</section></main></body></html>'''


def make_docx(talk: dict, spoken_chars: int, target: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    document.core_properties.title = talk["title"]
    document.core_properties.subject = "领导汇报：12 分钟主讲 + 3 分钟按需问答"
    document.core_properties.author = "PRISM 项目组"
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.3)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor.from_string("17202B")
    document.styles["Normal"].font.size = Pt(11.5)
    document.styles["Normal"].paragraph_format.line_spacing = 1.55
    document.styles["Normal"].paragraph_format.space_after = Pt(10)
    document.styles["Title"].font.size = Pt(27)
    document.styles["Heading 1"].font.size = Pt(19)
    document.styles["Heading 2"].font.size = Pt(14)
    header = section.header.paragraphs[0]
    header.text = "PRISM 棱镜 / 领导汇报讲稿 / 2026.10.26"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string("566173")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("PRISM · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.add_heading(talk["title"], 0)
    document.add_paragraph("PRISM 棱镜｜领导汇报演讲稿", "Subtitle")
    document.add_paragraph(f"{talk['mainCount']} 页主讲，约 12 分钟；4 页独立备用，预留 3 分钟按需问答。")
    document.add_paragraph(f"主讲正文与转场约 {spoken_chars:,} 字符。按每分钟约 220–240 字符排练，主讲包含约 15 秒猜图互动，并留看图与停顿时间；实际时长以试讲校准。只读口播正文和转场，不朗读标题、页面答案、讲者提示及来源。")
    document.add_heading("使用方式", 2)
    document.add_paragraph(f"第 2 页先讲三路用研，再看衰减曲线；第 3 页上下对照愿景；第 4 页讲两类难点与两个故事；第 5 页强调 MVP 完成范围。第 6 页猜图，留 15 秒互动；第 7 页揭晓 A 为 AI、B 为设计师，并介绍 Camon 进展；第 8 页从左向右讲生成流程。第 {talk['mainCount']} 页结束主讲，第 {talk['mainCount'] + 1}–{len(talk['slides'])} 页为 B1–B4 备答，每题约 40 秒，不全部连讲。猜图时不投屏讲稿或目录。讲稿弹层会显示在当前投屏画面，正式汇报请单独使用本文档。")
    document.add_heading("30 秒核心表达", 2)
    document.add_paragraph(talk["elevator"])
    for index, (slide, timing) in enumerate(zip(talk["slides"], timings(talk))):
        document.add_page_break()
        key = label(index, talk["mainCount"])
        document.add_heading(f"{key}｜{slide['title']}", 1)
        document.add_paragraph(timing, "Subtitle")
        document.add_paragraph(f"页面答案：{slide['answer']}")
        cue = document.add_paragraph(f"讲者提示：{slide['cue']}")
        for run in cue.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor.from_string("455A8B")
        for paragraph in slide["paragraphs"]:
            document.add_paragraph(paragraph)
        if slide["bridge"]:
            bridge = document.add_paragraph(f"转场：{slide['bridge']}")
            bridge.runs[0].italic = True
            bridge.runs[0].font.color.rgb = RGBColor.from_string("3557B1")
    document.add_page_break()
    document.add_heading("来源与表达边界", 1)
    document.add_paragraph("供准备，不逐条宣读。")
    for item in talk["sources"]:
        document.add_heading(item["label"], 2)
        document.add_paragraph(item["detail"])
    document.save(target)


def workflow_fragments(workflow: list[dict]) -> dict[str, str]:
    """生成第二页的工业设计阶段与同轴曲线；第五页独立表达产品能力链路。"""
    process = []
    for i, phase in enumerate(workflow):
        n = i + 1
        title = html.escape(phase["title"])
        # 上半部呈现多部门调研；下方保留完整阶段轴，辅助标签仍可读取阶段细节。
        detail = html.escape("；".join(phase["detail"].splitlines()), quote=True)
        process.append(f'<article class="process-step" aria-label="{title}：{detail}"><span class="phase-index">{n:02}</span><h3>{title}</h3></article>')
    # X 坐标使用八列中心点；没有访谈资料的项目启动阶段不虚构初始信号。
    guides = "".join(f'<line class="chart-guide" x1="{90 + i * 180}" y1="17" x2="{90 + i * 180}" y2="190"/>' for i in range(8))
    ticks = "".join(f'<text class="curve-index" x="{90 + i * 180}" y="216" text-anchor="middle">{i + 1:02}</text>' for i in range(8))
    dots = "".join(f'<circle class="signal-dot" cx="{x}" cy="{y}" r="5"/>' for x, y in [(270, 48), (450, 90), (630, 130), (810, 143), (990, 170), (1170, 122), (1350, 134)])
    curve = (
        '<svg class="signal-chart" viewBox="0 0 1440 230" role="img" aria-label="信号从用户研究阶段开始；语境压缩和多方取舍后减弱，重新引入用户反馈后局部回升。定性示意，非实测。">'
        + guides
        + '<path class="signal-path" d="M270 48C330 48 390 90 450 90S570 130 630 130S750 143 810 143S930 170 990 170S1110 122 1170 122S1290 134 1350 134"/>'
        + dots + ticks
        + '<text x="270" y="27" text-anchor="middle" class="curve-callout">原话与情境</text>'
        '<text x="450" y="65" text-anchor="middle">语境压缩</text>'
        '<text x="660" y="107" text-anchor="middle">视觉转换</text>'
        '<text x="965" y="145" text-anchor="middle">多方取舍</text>'
        '<text x="1190" y="90" text-anchor="middle" class="curve-callout">用户反馈，局部回升</text>'
        '<text x="90" y="126" text-anchor="middle">研究尚未展开</text></svg>'
    )
    return {"__PROCESS__": "".join(process), "__CURVE__": curve}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--presentation-only", action="store_true", help="仅重建演示与清单，保留独立讲稿和综合包")
    options = parser.parse_args()
    source_dir = Path(__file__).resolve().parent
    delivery = source_dir.parent
    talk = json.loads((source_dir / "talk.json").read_text())
    main_count = talk["mainCount"]
    assert main_count == 9 and len(talk["slides"]) == 13
    assert sum(slide["seconds"] for slide in talk["slides"][:main_count]) == talk["targetSeconds"] == 720
    assert talk["targetSeconds"] + talk["qaSeconds"] == 900
    # 来源身份由用户明确指定；保留原始文件字节，A/B 在猜图与揭晓页保持一致。
    board_files = {
        "a": delivery / "assets/board-a.png", "b": delivery / "assets/board-b.jpg",
        "mvp": delivery / "assets/mvp-moodboard.jpg",
        "dna": delivery / "assets/dna-10607-result.png",
        "dnaimage": delivery / "assets/dna-10607-original.png",
        "dnarelations": delivery / "assets/dna-10607-relations.png",
    }
    image_roles = {"a": "AI 情绪板", "b": "设计师情绪板", "mvp": "MVP 情绪板示例", "dna": "真实 DNA 完整 ID 结果截图", "dnaimage": "原图弹层截图", "dnarelations": "ID 七组关系区块截图"}
    board_bytes = {key: path.read_bytes() for key, path in board_files.items()}
    board_uris = {
        key: f"data:image/{'png' if board_files[key].suffix == '.png' else 'jpeg'};base64," + base64.b64encode(raw).decode("ascii")
        for key, raw in board_bytes.items()
    }
    slide_html = (source_dir / "slides-stage.html").read_text()
    footnotes = {
        1: "用户期待贯穿手机产品设计。右侧手机为形态示意。",
        2: "多部门调研共同服务下一代产品；设计阶段为流程初稿，曲线为定性风险示意。",
        3: "完整项目愿景：持续保留需求依据与影响力。曲线无数值，不代表零衰减实测。",
        4: "DNA 截图来自既有记录，具体状态可放大查看；专业维度、项目路径与两个故事依据项目方提供。",
        5: "一期聚焦情绪板聚类生成；二期延伸设计落地。规模、开发与产出物验证状态依据项目方提供。",
        6: "现场互动 · 原图完整展示、点击放大。两图展示不作为同条件受控实验。",
        7: "A＝AI，B＝设计师。Camon 调研及对比测试已开展，用户偏好结论尚未在此公布。",
        8: "数量与周期为项目汇报口径；3—5 天指按需选向、聚类与筛调，不包含历史资料积累及后续调研。",
        9: "预期效果 · 需求有设计响应，反馈能落实，选择理由持续保留。",
        10: "AI 提供候选与依据；设计师保留专业判断与修订入口。",
        11: "产出物验证与目标用户的偏好验证，分别回答不同问题。",
        12: "建议的采用口径 · 以实际设计决定及落实记录为依据。",
        13: "回到原始依据，保留未知，由人核对、选择和修订。",
    }
    for index, slide in enumerate(talk["slides"]):
        n = index + 1
        backup = index >= main_count
        key = label(index, main_count)
        chapter = f"备用 {key} / 按需展示" if backup else f"主讲 {key} / {main_count}"
        tag = "h1" if index == 0 else "h2"
        # 封面在逗号后换行，让产品价值表达完整地落在两行。
        question = html.escape(slide["title"])
        if index == 0:
            question = question.replace("棱镜：", '<span class="cover-name">棱镜：</span>', 1).replace("，", "，<br>", 1)
        frame = (
            f'<header class="slide-header"><div class="kicker"><span class="wordmark">PRISM</span><span class="divider"></span>'
            f'{"领导关心的问题 · 备用回答" if backup else html.escape(slide["part"])}'
            f'<span class="chapter">{chapter}</span></div><{tag} class="slide-question">{question}</{tag}>'
            f'<p class="answer">{html.escape(slide["answer"])}</p></header>'
        )
        footer = f'<footer class="page-meta"><span>{footnotes.get(n, talk["title"])}</span><span>{chapter}</span></footer>'
        slide_html = slide_html.replace(f"__FRAME_{n}__", frame).replace(f"__FOOTER_{n}__", footer)
    assert "__FRAME_" not in slide_html and "__FOOTER_" not in slide_html
    workflow = json.loads((source_dir / "workflow.json").read_text())
    assert len(workflow) == 8
    for marker, content in workflow_fragments(workflow).items():
        slide_html = slide_html.replace(marker, content)
    assert not re.search(r"__[A-Z_]+__", slide_html)
    # 在构建期绑定原图：本地浏览器即使禁用、拦截或中断 JS，也能看到完整图片。
    # A/B 在两页重复内嵌的体积开销，换取单文件离线可读，不依赖运行时模板搬运。
    slide_html = re.sub(
        r'<img data-board="([a-z]+)"',
        lambda match: f'<img data-board="{match[1]}" src="{board_uris[match[1]]}"',
        slide_html,
    )
    page = (source_dir / "roadshow.template.html").read_text()
    replacements = {
        "__STYLE__": "\n".join((source_dir / name).read_text() for name in ("stage.css", "generation-horizontal.css", "startup.css")),
        "__MAIN_COUNT__": str(main_count),
        "__TOTAL_COUNT__": str(len(talk["slides"])),
        "__SLIDES__": slide_html,
        "__TALK_JSON__": json.dumps(talk, ensure_ascii=False).replace("<", "\\u003c"),
        "__VIEWER__": (source_dir / "viewer.js").read_text(),
    }
    for marker, content in replacements.items():
        page = page.replace(marker, content)
    spoken = "".join("".join(slide["paragraphs"]) + slide["bridge"] for slide in talk["slides"][:main_count])
    spoken_chars = len(re.findall(r"[\u4e00-\u9fffA-Za-z0-9]", spoken))
    (delivery / "index.html").write_text(page, encoding="utf-8")
    # 修复演示交互不应覆写用户在讲稿中的手工调整；内容更新任务才完整重建讲稿。
    if not options.presentation_only:
        markdown = make_markdown(talk, spoken_chars)
        for name in ("Untitled.md", "演讲稿.md"):
            (delivery / name).write_text(markdown, encoding="utf-8")
        (delivery / "演讲稿.html").write_text(make_speech_html(talk, spoken_chars), encoding="utf-8")
        make_docx(talk, spoken_chars, delivery / "PRISM-产品路演演讲稿.docx")
    manifest = {
        "title": talk["title"], "date": talk["date"], "main_slide_count": main_count,
        "backup_slide_count": len(talk["slides"]) - main_count, "slide_count": len(talk["slides"]),
        "talk_seconds": talk["targetSeconds"], "qa_seconds": talk["qaSeconds"],
        "spoken_character_count": spoken_chars,
        "interaction_seconds": talk["interactionSeconds"],
        "prior_version": "history/before-speech-sync-20261026.zip",
        "updated_scope": "HTML 与内嵌讲者提示；独立讲稿及综合 ZIP 本次未重建，PPT 保留旧版" if options.presentation_only else "HTML 内嵌讲稿、独立 Markdown／HTML／Word 稿及综合 ZIP 同步为 9 页主讲＋4 页备答；演示视觉未变，截图 PPT 保留旧版且不入包",
        "images": [
            {"slot": key.upper(), "role": image_roles[key],
             "filename": str(board_files[key].relative_to(delivery)),
             "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw),
             "origin": "用户提供原图" if not key.startswith("dna") else "https://gamedevcenter.ahagamecenter.com/dna-v2 的现有记录截图，详见 assets/dna-evidence.json",
             "pages": [6, 7] if key in ("a", "b") else [8 if key == "mvp" else 4]}
            for key, raw in board_bytes.items()
        ],
        "files": {},
    }
    filenames = ["index.html", "Untitled.md", "PRISM-产品路演演讲稿.docx", "演讲稿.html", "演讲稿.md", "README.md", "assets/dna-evidence.json"]
    for filename in filenames:
        artifact = delivery / filename
        manifest["files"][filename] = {"bytes": artifact.stat().st_size, "sha256": hashlib.sha256(artifact.read_bytes()).hexdigest()}
    (delivery / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    # 用户仅要求 HTML 时，不用新版页面与旧版讲稿重打综合包，避免交付版本混淆。
    if not options.presentation_only:
        with zipfile.ZipFile(delivery / "PRISM-产品路演与演讲稿.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for filename in [*filenames, "manifest.json"]:
                archive.write(delivery / filename, filename)
    print(json.dumps({"main": main_count, "backup": 4, "main_seconds": 720, "qa_seconds": 180, "main_chars": spoken_chars}, ensure_ascii=False))


if __name__ == "__main__":
    main()
