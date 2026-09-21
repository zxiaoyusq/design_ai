"""从原始文章表生成前端静态快照，图片复用已有下载文件。"""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "ref/文章信息表.xlsx"
IMAGE_INDEX = ROOT / "data/trend_data/trends.json"
OUTPUT = ROOT / "frontend/src/data/articleTrends.ts"
LIMIT = 30


def calendar_date(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value or "").strip()[:10]


def js_string(value: object) -> str:
    return json.dumps(value, ensure_ascii=False)


def main() -> None:
    image_index = json.loads(IMAGE_INDEX.read_text(encoding="utf-8"))
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    if image_index["source"]["sha256"] != source_hash:
        raise ValueError("图片索引对应的源表与当前文章信息表不同，请先核对图片归属")

    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook[image_index["source"]["sheet"]]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    articles = []
    for row_number, values in enumerate(rows, start=2):
        article = dict(zip(headers, values))
        article["source_row"] = row_number
        article["release_time"] = calendar_date(article["release_time"])
        articles.append(article)
    workbook.close()

    # 同日文章沿用表格顺序；无发布日期的行不会占用最近 30 条名额。
    recent = sorted(articles, key=lambda item: (item["release_time"], -item["source_row"]), reverse=True)[:LIMIT]
    indexed = {article["id"]: article for article in image_index["trends"]}
    lines = [
        "// 由 frontend/scripts/prepare_article_trends.py 从 ref/文章信息表.xlsx 生成。",
        f"// 源表 SHA-256: {source_hash}；按发布日期降序，同日按表格行序取最近 {LIMIT} 条。",
        "import type { ArticleTrend } from '@/types/articleTrend'",
        "",
        "export const articleTrends: ArticleTrend[] = [",
    ]
    for article in recent:
        article_id = str(article["id"])
        indexed_article = indexed[article_id]
        urls = [url.strip() for url in str(article["image_url"] or "").split("||") if url.strip()]
        images = indexed_article["images"]
        if [image["url"] for image in images] != urls:
            raise ValueError(f"文章 {article_id} 的图片顺序与源表不符")
        lines.append("  {")
        for field, source_field in (
            ("id", "id"),
            ("sourceRow", "source_row"),
            ("title", "title_zh"),
            ("originalTitle", "title"),
            ("summary", "summary_zh"),
            ("sourceUrl", "source_url"),
            ("detailUrl", "detail_url"),
            ("releaseDate", "release_time"),
            ("category", "primary_category"),
        ):
            lines.append(f"    {field}: {js_string(article[source_field])},")
        for field in ("subcategory", "tags"):
            value = article[field]
            items = [part.strip() for part in str(value or "").replace("，", ",").split(",") if part.strip()]
            lines.append(f"    {field if field == 'tags' else 'subcategories'}: {js_string(items)},")
        lines.append("    images: [")
        for image in images:
            local_path = image["local_path"]
            if image["status"] != "downloaded" or not local_path:
                continue
            path = (IMAGE_INDEX.parent / local_path).resolve()
            if not path.is_relative_to(IMAGE_INDEX.parent.resolve()) or not path.is_file():
                raise ValueError(f"文章 {article_id} 的图片路径无效: {local_path}")
            # 静态 URL 会让 Vite 仅将这 30 篇文章用到的本地图片打包。
            lines.append(f"      new URL({js_string('../../../data/trend_data/' + local_path)}, import.meta.url).href,")
        lines.extend(["    ],", "  },"])
    lines.append("]")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"已生成 {len(recent)} 条文章：{OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
